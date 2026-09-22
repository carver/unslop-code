"""Checks for the export / upload / multi-format specification."""
import csv
import io
import os
import socket
import subprocess
import sys
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer

import openpyxl
import requests
import xlwt

HERE = os.path.dirname(os.path.abspath(__file__))

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS = "application/vnd.ms-excel"


def make_xlsx(sheets):
    """sheets: [(title, [[cell, ...], ...]), ...] -> .xlsx bytes."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title=title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_xls(sheets):
    """sheets: [(title, [[cell, ...], ...]), ...] -> legacy .xls bytes."""
    book = xlwt.Workbook()
    for title, rows in sheets:
        sheet = book.add_sheet(title)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value is not None:
                    sheet.write(r, c, value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


BASIC_ROWS = [["name", "age", "score", "start"],
              ["Alice", 30, 95.5, "08:30"],
              ["Bob", 7, -2.25, "9:15"]]

XLSX_BASIC = make_xlsx([("Data", BASIC_ROWS)])
XLS_BASIC = make_xls([("Data", BASIC_ROWS)])
XLSX_TWO = make_xlsx([("First", [["a", "b"], [1, 2]]),
                      ("Second", [["x", "y", "z"], [9, 9, 9]])])
XLS_TWO = make_xls([("First", [["a", "b"], [1, 2]]),
                    ("Second", [["x", "y", "z"], [9, 9, 9]])])
XLSX_HEADER_ONLY = make_xlsx([("First", [["a", "b"]]),
                              ("Second", [["x", "y"], [1, 2]])])
XLSX_EMPTY = make_xlsx([("First", [])])

_plain_zip = io.BytesIO()
with zipfile.ZipFile(_plain_zip, "w") as _z:
    _z.writestr("readme.txt", "hello")
PLAIN_ZIP = _plain_zip.getvalue()

CSV_BASIC = b"name,age,score,start\nAlice,30,95.5,08:30\nBob,7,-2.25,9:15\n"
CSV_LATIN1 = "naive,ville\nr\xe9sum\xe9,Gen\xe8ve\n".encode("latin-1")
CSV_QUOTED = b'name,note\n"Smith, John","says ""hi"", ok"\n"multi\nline",2\n'
CSV_BIG = ("n\n" + "".join("%d\n" % i for i in range(500))).encode()
CSV_SORT = b"k,v\nb,1\na,2\nc,3\n"

FIXTURES = {
    "/basic.csv": (CSV_BASIC, "text/csv"),
    "/basic.xlsx": (XLSX_BASIC, XLSX),
    "/basic.xls": (XLS_BASIC, XLS),
    "/two.xlsx": (XLSX_TWO, XLSX),
    "/two.xls": (XLS_TWO, XLS),
    "/headeronly.xlsx": (XLSX_HEADER_ONLY, XLSX),
    "/empty.xlsx": (XLSX_EMPTY, XLSX),
    "/plain.zip": (PLAIN_ZIP, "application/zip"),
    "/doc.pdf": (b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\ntrailer\n", "application/pdf"),
    "/latin1.csv": (CSV_LATIN1, "text/csv"),
}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        path = self.path.split("?")[0]
        if path not in FIXTURES:
            self.send_error(404)
            return
        body, ctype = FIXTURES[path]
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


FAILURES = []
CHECKS = [0]


def check(name, cond, detail=""):
    CHECKS[0] += 1
    if cond:
        print("  PASS  %s" % name)
    else:
        print("  FAIL  %s %s" % (name, detail))
        FAILURES.append(name)


def parse_csv(text):
    return [row for row in csv.reader(io.StringIO(text, newline=""))]


def main():
    fixture_port = free_port()
    httpd = HTTPServer(("127.0.0.1", fixture_port), Handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    FIX = "http://127.0.0.1:%d" % fixture_port

    app_port = free_port()
    proc = subprocess.Popen(
        [sys.executable, os.path.join(HERE, "datagate.py"), "start",
         "--port", str(app_port), "--address", "127.0.0.1"],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
    )
    BASE = "http://127.0.0.1:%d" % app_port
    for _ in range(100):
        try:
            requests.get(BASE + "/", timeout=1)
            break
        except Exception:
            time.sleep(0.1)
    else:
        out = proc.stdout.read().decode() if proc.stdout else ""
        raise SystemExit("server never started\n" + out)

    def conv(path, **kw):
        params = {"source": FIX + path}
        params.update(kw)
        return requests.get(BASE + "/convert", params=params, timeout=20)

    def up(data, field="file", filename="data.csv", **kw):
        return requests.post(BASE + "/upload", params=kw,
                             files={field: (filename, data)}, timeout=20)

    def ds(endpoint, **kw):
        return requests.get(BASE + endpoint, params=kw, timeout=20)

    def exp(endpoint, **kw):
        return requests.get(BASE + endpoint + "/export", params=kw, timeout=20)

    def envelope(r, status):
        try:
            j = r.json()
        except Exception:
            return False
        return r.status_code == status and j.get("ok") is False and "error" in j

    try:
        # ================= EXPORT =================
        e = conv("/basic.csv").json()["endpoint"]
        ds_id = e.rsplit("/", 1)[1]
        r = exp(e)
        check("export 200", r.status_code == 200, r.text)
        check("export content-type", r.headers.get("Content-Type") == "text/csv",
              r.headers.get("Content-Type"))
        check("export content-disposition",
              r.headers.get("Content-Disposition")
              == 'attachment; filename="%s.csv"' % ds_id,
              r.headers.get("Content-Disposition"))
        rows = parse_csv(r.content.decode("utf-8"))
        check("export header is source column order",
              rows[0] == ["name", "age", "score", "start"], rows)
        check("export rows", rows[1:] == [["Alice", "30", "95.5", "08:30"],
                                          ["Bob", "7", "-2.25", "9:15"]], rows)

        # filters / sort / pagination mirror /datasets/<id>
        rows = parse_csv(exp(e, **{"name__exact": "Bob"}).content.decode())
        check("export applies filter", rows == [["name", "age", "score", "start"],
                                                ["Bob", "7", "-2.25", "9:15"]], rows)
        rows = parse_csv(exp(e, **{"age__greater": "100"}).content.decode())
        check("export filter to empty", rows == [["name", "age", "score", "start"]], rows)

        e_sort = conv("/basic.csv").json()["endpoint"]
        rows = parse_csv(exp(e_sort, _sort="age").content.decode())
        check("export applies _sort", [r[0] for r in rows[1:]] == ["Bob", "Alice"], rows)
        rows = parse_csv(exp(e_sort, _sort_desc="age").content.decode())
        check("export applies _sort_desc", [r[0] for r in rows[1:]] == ["Alice", "Bob"], rows)

        rows = parse_csv(exp(e, _size="1").content.decode())
        check("export applies _size", rows == [["name", "age", "score", "start"],
                                               ["Alice", "30", "95.5", "08:30"]], rows)
        rows = parse_csv(exp(e, _size="1", _offset="1").content.decode())
        check("export applies _offset", rows == [["name", "age", "score", "start"],
                                                 ["Bob", "7", "-2.25", "9:15"]], rows)

        e_big = up(CSV_BIG, filename="big.csv").json()["endpoint"]
        rows = parse_csv(exp(e_big).content.decode())
        check("export default page size", len(rows) == 101, len(rows))
        rows = parse_csv(exp(e_big, _size="5", _offset="2").content.decode())
        check("export paginates big", rows == [["n"], ["2"], ["3"], ["4"], ["5"], ["6"]], rows)

        # shape/rowid/total do not affect CSV
        base = exp(e).content
        for name, value in [("_shape", "objects"), ("_shape", "lists"),
                            ("_rowid", "hide"), ("_total", "hide")]:
            r2 = exp(e, **{name: value})
            check("export ignores %s=%s" % (name, value),
                  r2.status_code == 200 and r2.content == base, r2.text[:200])
        r2 = exp(e, _shape="objects", _rowid="hide", _total="hide")
        check("export ignores all three", r2.content == base, r2.text[:200])

        # quoting round-trips
        e_q = up(CSV_QUOTED, filename="q.csv").json()["endpoint"]
        want = ds(e_q).json()["rows"]
        got = parse_csv(exp(e_q).content.decode())
        check("export quoting round-trips",
              got[0] == ["name", "note"]
              and [[str(c) for c in row] for row in want] == got[1:], (want, got))

        # errors
        check("export unknown dataset 404", envelope(exp("/datasets/deadbeef"), 404))
        check("export bad _size 400", envelope(exp(e, _size="0"), 400))
        check("export bad _sort 400", envelope(exp(e, _sort="nope"), 400))
        check("export unknown filter column 400",
              envelope(exp(e, **{"nope__exact": "x"}), 400))

        # ================= UPLOAD =================
        r = up(CSV_BASIC)
        j = r.json()
        check("upload 200", r.status_code == 200, r.text)
        check("upload ok true", j.get("ok") is True, j)
        check("upload endpoint shape",
              isinstance(j.get("endpoint"), str)
              and j["endpoint"].startswith("/datasets/"), j)
        d = ds(j["endpoint"]).json()
        check("upload dataset queryable",
              d["columns"] == ["name", "age", "score", "start"]
              and d["rows"] == [["Alice", 30, 95.5, "08:30"], ["Bob", 7, -2.25, "9:15"]], d)

        r2 = up(CSV_BASIC, field="attachment", filename="other.csv")
        check("upload attachment field 200", r2.status_code == 200, r2.text)
        check("same bytes -> same id", r2.json()["endpoint"] == j["endpoint"],
              (r2.json(), j))
        r3 = up(CSV_BASIC, filename="totally-different-name.csv")
        check("same bytes different name -> same id",
              r3.json()["endpoint"] == j["endpoint"])
        r4 = up(b"a,b\n1,2\n")
        check("different bytes -> different id", r4.json()["endpoint"] != j["endpoint"])

        # non-multipart -> 415
        r = requests.post(BASE + "/upload", data=CSV_BASIC,
                          headers={"Content-Type": "text/csv"}, timeout=10)
        check("non-multipart 415", envelope(r, 415), r.text[:200])
        r = requests.post(BASE + "/upload", json={"file": "a,b\n1,2"}, timeout=10)
        check("json body 415", envelope(r, 415), r.text[:200])
        r = requests.post(BASE + "/upload", data=b"", timeout=10)
        check("no content-type 415", envelope(r, 415), r.text[:200])
        r = requests.post(BASE + "/upload", data={"x": "1"}, timeout=10)
        check("urlencoded form 415", envelope(r, 415), r.text[:200])

        # missing field / malformed -> 400
        r = requests.post(BASE + "/upload", files={"other": ("a.csv", CSV_BASIC)},
                          timeout=10)
        check("missing file field 400", envelope(r, 400), r.text[:200])
        r = requests.post(
            BASE + "/upload", data=b"this is not a multipart body at all",
            headers={"Content-Type": "multipart/form-data; boundary=xyzzy"}, timeout=10)
        check("malformed multipart 400", envelope(r, 400), r.text[:200])
        r = requests.post(BASE + "/upload", data=b"--x\r\ngarbage",
                          headers={"Content-Type": "multipart/form-data"}, timeout=10)
        check("multipart without boundary 400", envelope(r, 400), r.text[:200])

        # charset on upload
        r = up(CSV_LATIN1, filename="l.csv", charset="latin-1")
        check("upload charset applied",
              ds(r.json()["endpoint"]).json()["rows"] == [["r\xe9sum\xe9", "Gen\xe8ve"]],
              ds(r.json()["endpoint"]).json())
        check("upload bad charset 400",
              envelope(up(CSV_BASIC, charset="not-a-charset"), 400))
        check("upload undecodable charset 400",
              envelope(up(CSV_LATIN1, filename="l.csv", charset="utf-8"), 400))

        # non-tabular / unsupported uploads
        check("upload empty file 400", envelope(up(b"", filename="e.csv"), 400))
        check("upload html 400",
              envelope(up(b"<html><body>hi</body></html>", filename="p.html"), 400))
        check("upload pdf 400", envelope(up(b"%PDF-1.4\nstuff\n", filename="d.pdf"), 400))
        check("upload plain zip 400", envelope(up(PLAIN_ZIP, filename="a.zip"), 400))
        check("upload png 400",
              envelope(up(b"\x89PNG\r\n\x1a\n" + bytes(range(256)), filename="i.png"), 400))

        check("GET /upload 405", envelope(requests.get(BASE + "/upload", timeout=10), 405))

        # ================= MULTI-FORMAT =================
        for label, blob, name in [("xlsx", XLSX_BASIC, "b.xlsx"), ("xls", XLS_BASIC, "b.xls")]:
            r = up(blob, filename=name)
            check("upload %s 200" % label, r.status_code == 200, r.text[:200])
            d = ds(r.json()["endpoint"]).json()
            check("upload %s columns in source order" % label,
                  d["columns"] == ["name", "age", "score", "start"], d)
            check("upload %s rows" % label,
                  d["rows"] == [["Alice", 30, 95.5, "08:30"],
                                ["Bob", 7, -2.25, "9:15"]], d)

        for label, path in [("xlsx", "/basic.xlsx"), ("xls", "/basic.xls")]:
            r = conv(path)
            check("convert %s 200" % label, r.status_code == 200, r.text[:200])
            d = ds(r.json()["endpoint"]).json()
            check("convert %s parsed" % label,
                  d["columns"] == ["name", "age", "score", "start"]
                  and d["rows"][0] == ["Alice", 30, 95.5, "08:30"], d)

        # only the first worksheet
        for label, blob, name in [("xlsx", XLSX_TWO, "t.xlsx"), ("xls", XLS_TWO, "t.xls")]:
            d = ds(up(blob, filename=name).json()["endpoint"]).json()
            check("%s first worksheet only" % label,
                  d["columns"] == ["a", "b"] and d["rows"] == [[1, 2]], d)
        for label, path in [("xlsx", "/two.xlsx"), ("xls", "/two.xls")]:
            d = ds(conv(path).json()["endpoint"]).json()
            check("convert %s first worksheet only" % label,
                  d["columns"] == ["a", "b"] and d["rows"] == [[1, 2]], d)

        # first sheet must be tabular
        check("xlsx header-only first sheet 400",
              envelope(up(XLSX_HEADER_ONLY, filename="h.xlsx"), 400))
        check("convert xlsx header-only 400", envelope(conv("/headeronly.xlsx"), 400))
        check("xlsx empty first sheet 400",
              envelope(up(XLSX_EMPTY, filename="e.xlsx"), 400))

        # charset validates only for text CSV
        r = up(XLSX_BASIC, filename="b.xlsx", charset="not-a-charset")
        check("charset ignored for xlsx", r.status_code == 200, r.text[:200])
        r = up(XLS_BASIC, filename="b.xls", charset="not-a-charset")
        check("charset ignored for xls", r.status_code == 200, r.text[:200])
        r = conv("/basic.xlsx", charset="bogus-codec")
        check("convert charset ignored for xlsx", r.status_code == 200, r.text[:200])
        check("convert bad charset still 400 for csv",
              envelope(conv("/basic.csv", charset="not-a-charset"), 400))

        # unrecognized formats
        check("convert plain zip 400", envelope(conv("/plain.zip"), 400))
        check("convert pdf 400", envelope(conv("/doc.pdf"), 400))

        # content decides the format, not the file name
        d = ds(up(XLSX_BASIC, filename="lies.csv").json()["endpoint"]).json()
        check("xlsx bytes named .csv still parse", d["columns"] == ["name", "age", "score", "start"], d)
        d = ds(up(CSV_BASIC, filename="lies.xlsx").json()["endpoint"]).json()
        check("csv bytes named .xlsx still parse", d["columns"] == ["name", "age", "score", "start"], d)

        # spreadsheets export as CSV too
        e_x = up(XLSX_BASIC, filename="b.xlsx").json()["endpoint"]
        r = exp(e_x)
        check("export spreadsheet content-type", r.headers.get("Content-Type") == "text/csv")
        check("export spreadsheet rows",
              parse_csv(r.content.decode()) == [["name", "age", "score", "start"],
                                                ["Alice", "30", "95.5", "08:30"],
                                                ["Bob", "7", "-2.25", "9:15"]],
              r.content)
        rows = parse_csv(exp(e_x, **{"name__exact": "Bob"}, _size="1").content.decode())
        check("export spreadsheet filter+size",
              rows == [["name", "age", "score", "start"], ["Bob", "7", "-2.25", "9:15"]], rows)

        # filters/sort work on spreadsheet datasets
        d = ds(e_x, **{"age__greater": "10"}).json()
        check("spreadsheet numeric filter", d["rows"] == [["Alice", 30, 95.5, "08:30"]], d)
        d = ds(e_x, _sort="age").json()
        check("spreadsheet sort", [r[0] for r in d["rows"]] == ["Bob", "Alice"], d)

        # unicode survives export
        e_u = up(CSV_LATIN1, filename="l.csv", charset="latin-1").json()["endpoint"]
        got = parse_csv(exp(e_u).content.decode("utf-8"))
        check("export unicode", got == [["naive", "ville"], ["r\xe9sum\xe9", "Gen\xe8ve"]], got)

        # determinism: repeated export is byte-identical
        check("export deterministic", exp(e).content == exp(e).content)

        # ================= PROTOCOL DETAILS =================
        r = requests.head(BASE + e + "/export", timeout=10)
        check("HEAD export headers",
              r.status_code == 200
              and r.headers.get("Content-Type") == "text/csv"
              and r.headers.get("Content-Disposition")
              == 'attachment; filename="%s.csv"' % ds_id, dict(r.headers))
        check("export sends CORS",
              exp(e).headers.get("Access-Control-Allow-Origin") is not None)
        r = requests.options(BASE + "/upload", timeout=10,
                             headers={"Origin": "http://x.test",
                                      "Access-Control-Request-Method": "POST"})
        check("upload preflight", r.status_code in (200, 204)
              and r.headers.get("Access-Control-Allow-Origin") is not None, r.status_code)
        r = requests.options(BASE + e + "/export", timeout=10,
                             headers={"Origin": "http://x.test",
                                      "Access-Control-Request-Method": "GET"})
        check("export preflight", r.status_code in (200, 204), r.status_code)
        check("export bad _shape 400", envelope(exp(e, _shape="bogus"), 400))
        check("export bad _rowid 400", envelope(exp(e, _rowid="yes"), 400))
        check("export repeated control 400",
              envelope(requests.get(BASE + e + "/export?_size=1&_size=2", timeout=10), 400))

        # a part sent without a filename is still the file field
        r = requests.post(BASE + "/upload", files={"file": (None, CSV_BASIC.decode())},
                          timeout=10)
        check("upload plain form part 200", r.status_code == 200, r.text[:200])
        check("upload plain form part same id", r.json()["endpoint"] == j["endpoint"],
              r.json())
        r = requests.post(BASE + "/upload",
                          files={"attachment": (None, CSV_BASIC.decode())}, timeout=10)
        check("upload plain attachment part 200", r.status_code == 200, r.text[:200])

        # charset does not change the id of identical bytes
        a = up(CSV_BASIC, charset="utf-8").json()["endpoint"]
        check("charset does not change id", a == j["endpoint"], (a, j))

        # when both parts are present, file wins
        r = requests.post(BASE + "/upload",
                          files=[("file", ("a.csv", CSV_BASIC)),
                                 ("attachment", ("b.csv", b"z\n1\n"))], timeout=10)
        check("file part preferred over attachment",
              r.json().get("endpoint") == j["endpoint"], r.text[:200])

        check("POST /convert 405",
              envelope(requests.post(BASE + "/convert", timeout=10), 405))
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=5)
        except Exception:
            proc.kill()
        httpd.shutdown()

    print("\n%d checks, %d failures" % (CHECKS[0], len(FAILURES)))
    if FAILURES:
        print("failed: " + ", ".join(FAILURES))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
