"""Builders for the binary workbook payloads the multi-format tests ingest."""

import io

import openpyxl
import xlwt


def make_xlsx(sheets):
    """Serialise `sheets`, a list of `(title, rows)` pairs, as `.xlsx` bytes."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_xls(sheets):
    """Serialise `sheets`, a list of `(title, rows)` pairs, as legacy `.xls` bytes."""
    book = xlwt.Workbook()
    for title, rows in sheets:
        sheet = book.add_sheet(title)
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                sheet.write(row_index, column_index, value)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
