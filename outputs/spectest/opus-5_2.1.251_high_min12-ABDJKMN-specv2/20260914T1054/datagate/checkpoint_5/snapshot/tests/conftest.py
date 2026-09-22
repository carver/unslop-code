"""Shared fixtures: a real local origin server for `source` URLs, plus a client for the app.

The spec's `/convert` fetches a *remote* CSV over HTTP, so the tests stand up a throwaway
HTTP origin on 127.0.0.1 and hand out real URLs into it. No outbound network is used.
"""
import csv
import http.server
import io
import socket
import threading
import zipfile

import pytest

import datagate


class _OriginHandler(http.server.BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, *args):  # keep the test output clean, no undrained pipes
        pass

    def _route(self):
        path = self.path.split("?")[0]
        self.server.hits[path] = self.server.hits.get(path, 0) + 1
        return self.server.routes.get(path)

    def do_GET(self):
        route = self._route()
        if route is None:
            body = b"not found"
            self.send_response(404)
            self.send_header("Content-Type", "text/plain")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        status, content_type, body = route
        self.send_response(status)
        if content_type is not None:
            self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_HEAD(self):
        self.do_GET()


class Origin:
    """Serves registered byte payloads and vends the URL for each."""

    def __init__(self):
        self._server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _OriginHandler)
        self._server.routes = {}
        self._server.hits = {}
        self._server.daemon_threads = True
        self.port = self._server.server_address[1]
        self.base = "http://127.0.0.1:%d" % self.port
        self._thread = threading.Thread(target=self._server.serve_forever, daemon=True)
        self._thread.start()
        self._n = 0

    def add(self, body, path=None, status=200, content_type="text/csv"):
        """Register a payload; returns the absolute URL that serves it."""
        if isinstance(body, str):
            body = body.encode("utf-8")
        if path is None:
            self._n += 1
            path = "/f%d.csv" % self._n
        self._server.routes[path] = (status, content_type, body)
        return self.base + path

    def update(self, url, body, status=200, content_type="text/csv"):
        """Change what an already-registered URL serves (same URL, new content)."""
        if isinstance(body, str):
            body = body.encode("utf-8")
        self._server.routes[self._path(url)] = (status, content_type, body)
        return url

    def remove(self, url):
        """Make an already-registered URL start answering 404."""
        self._server.routes.pop(self._path(url), None)
        return url

    def hits(self, url):
        """How many HTTP requests this origin has served for `url` so far."""
        return self._server.hits.get(self._path(url), 0)

    def _path(self, url):
        return url[len(self.base):].split("?")[0] if url.startswith(self.base) else url

    def url_for(self, path):
        return self.base + path

    def shutdown(self):
        self._server.shutdown()
        self._server.server_close()


@pytest.fixture(scope="session")
def origin():
    o = Origin()
    yield o
    o.shutdown()


@pytest.fixture()
def app():
    application = datagate.create_app()
    application.config.update(TESTING=True)
    return application


@pytest.fixture()
def client(app):
    return app.test_client()


@pytest.fixture()
def convert(client):
    """POST-free helper: GET /convert with query params, returning (status, json)."""

    def _convert(**params):
        query = {k: v for k, v in params.items() if v is not None}
        resp = client.get("/convert", query_string=query)
        return resp.status_code, resp.get_json()

    return _convert


@pytest.fixture()
def dataset(client, convert, origin):
    """Ingest a CSV body and return the parsed /datasets/<id> payload."""

    def _dataset(body, charset=None, query_string=None, content_type="text/csv"):
        url = origin.add(body, content_type=content_type)
        status, payload = convert(source=url, charset=charset)
        assert status == 200, payload
        resp = client.get(payload["endpoint"], query_string=query_string or {})
        return resp.status_code, resp.get_json()

    return _dataset


def unused_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


# ------------------------------------------ multi-format / upload helpers ---

def xlsx_bytes(rows, sheet_name="Sheet1", extra_sheets=()):
    """Build a real .xlsx workbook; `rows` is a list of row lists."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(list(row))
    for name, other_rows in extra_sheets:
        other = book.create_sheet(title=name)
        for row in other_rows:
            other.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def xls_bytes(rows, sheet_name="Sheet1", extra_sheets=()):
    """Build a real (BIFF8) .xls workbook; `rows` is a list of row lists."""
    import xlwt

    book = xlwt.Workbook()

    def fill(sheet, data):
        for r, row in enumerate(data):
            for c, value in enumerate(row):
                if value is None:
                    continue
                sheet.write(r, c, value)

    fill(book.add_sheet(sheet_name), rows)
    for name, other_rows in extra_sheets:
        fill(book.add_sheet(name), other_rows)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def zip_bytes(names_to_data):
    """A plain (non-xlsx) zip archive - a recognisable container, wrong format."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in names_to_data.items():
            archive.writestr(name, data)
    return buffer.getvalue()


@pytest.fixture()
def upload(client):
    """POST /upload with a multipart body; returns (status, json-or-None)."""

    def _upload(body=b"name,qty\nwidget,3\n", field="file", filename="data.csv",
                charset=None, content_type="multipart/form-data", extra=None):
        if isinstance(body, str):
            body = body.encode("utf-8")
        data = {}
        if field is not None:
            data[field] = (io.BytesIO(body), filename)
        if extra:
            data.update(extra)
        query = {} if charset is None else {"charset": charset}
        resp = client.post(
            "/upload", data=data, content_type=content_type, query_string=query
        )
        return resp.status_code, resp.get_json(silent=True)

    return _upload


@pytest.fixture()
def uploaded(client, upload):
    """Upload bytes and return the parsed /datasets/<id> payload."""

    def _uploaded(body, query_string=None, **kwargs):
        status, payload = upload(body=body, **kwargs)
        assert status == 200, payload
        resp = client.get(payload["endpoint"], query_string=query_string or {})
        return resp.status_code, resp.get_json()

    return _uploaded


@pytest.fixture()
def export(client, convert, origin):
    """Ingest a CSV body, then GET its /export; returns the response object."""

    def _export(body, query_string=None, charset=None):
        url = origin.add(body)
        status, payload = convert(source=url, charset=charset)
        assert status == 200, payload
        return client.get(payload["endpoint"] + "/export",
                          query_string=query_string or {})

    return _export


def csv_rows(response):
    """Parse an export response body back into a list of row lists."""
    text = response.get_data(as_text=True)
    return [row for row in csv.reader(io.StringIO(text, newline="")) if row != []]
