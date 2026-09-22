"""End-to-end tests for the datagate HTTP API and CSV pipeline."""

import csv
import io
import pathlib
import zipfile

import openpyxl
import pytest
import requests
import xlwt

import config
from app import create_app
from config import ConfigError
from errors import ApiError
from tabular import decode, detect_delimiter, infer_value

SAMPLE_CSV = b"name,start,score,ratio\nAda,08:30,42,1.5\nGrace,9:15,7,0.25\n"


@pytest.fixture
def client(monkeypatch, tmp_path):
    """A test client whose fetches are served from an in-test payload table.

    Each test gets its own storage directory, so no run inherits the datasets
    of the one before it.
    """
    monkeypatch.setenv("STORAGE_DIR", str(tmp_path / "datasets"))
    payloads = {"https://example.com/data.csv": SAMPLE_CSV}
    fetched = []

    def fake_fetch(url):
        fetched.append(url)
        if url not in payloads:
            raise ApiError(f"Source unreachable: {url}", 404)
        return payloads[url]

    monkeypatch.setattr("app.fetch", fake_fetch)
    app = create_app()
    app.config["TESTING"] = True
    client = app.test_client()
    client.payloads = payloads
    client.fetched = fetched
    return client


def convert(client, **params):
    return client.get("/convert", query_string=params)


def big_dataset(client):
    """Convert a 250 row source and return its dataset endpoint."""
    rows = "\n".join(f"row{index};{index}" for index in range(250))
    client.payloads["https://example.com/big.csv"] = f"label;value\n{rows}\n".encode()
    return convert(client, source="https://example.com/big.csv").json["endpoint"]


def test_convert_returns_stable_endpoint(client):
    first = convert(client, source="https://example.com/data.csv")
    second = convert(client, source="https://example.com/data.csv")

    assert first.status_code == 200
    assert first.json["ok"] is True
    assert first.json["endpoint"] == second.json["endpoint"]


def test_dataset_preserves_order_and_infers_types(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint)

    assert response.status_code == 200
    assert response.json["columns"] == ["name", "start", "score", "ratio"]
    assert response.json["rows"] == [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]
    assert response.json["total"] == 2
    assert response.json["query_ms"] >= 0


def test_rows_are_capped_at_one_hundred_by_default(client):
    endpoint = big_dataset(client)

    assert len(client.get(endpoint).json["rows"]) == 100
    assert len(client.get(endpoint, query_string={"_size": 5}).json["rows"]) == 5


def test_size_beyond_the_row_count_returns_everything(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert len(client.get(endpoint, query_string={"_size": 500}).json["rows"]) == 2


def test_offset_skips_rows_and_total_ignores_pagination(client):
    endpoint = big_dataset(client)

    response = client.get(endpoint, query_string={"_offset": 248, "_size": 10})

    assert response.json["rows"] == [["row248", 248], ["row249", 249]]
    assert response.json["total"] == 250


def test_offset_past_the_end_returns_no_rows(client):
    endpoint = big_dataset(client)

    assert client.get(endpoint, query_string={"_offset": 400}).json["rows"] == []


def test_total_counts_rows_before_pagination(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert client.get(endpoint, query_string={"_size": 1}).json["total"] == 2


def test_sorting_runs_before_pagination(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    ascending = client.get(endpoint, query_string={"_sort": "score", "_size": 1})
    descending = client.get(endpoint, query_string={"_sort_desc": "score", "_size": 1})

    assert ascending.json["rows"] == [["Grace", "9:15", 7, 0.25]]
    assert descending.json["rows"] == [["Ada", "08:30", 42, 1.5]]


def test_descending_sort_wins_over_ascending(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string={"_sort": "name", "_sort_desc": "name"})

    assert [row[0] for row in response.json["rows"]] == ["Grace", "Ada"]


def test_sorting_keeps_tied_rows_in_source_order(client):
    client.payloads["https://example.com/ties.csv"] = (
        b"name,team\nAda,blue\nGrace,blue\nKay,amber\n"
    )
    endpoint = convert(client, source="https://example.com/ties.csv").json["endpoint"]

    for params in ({"_sort": "team"}, {"_sort_desc": "team"}):
        response = client.get(endpoint, query_string=params)
        assert [row[0] for row in response.json["rows"] if row[1] == "blue"] == ["Ada", "Grace"]


def test_object_shape_carries_rowid(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string={"_shape": "objects", "_size": 1})

    assert response.json["rows"] == [
        {"rowid": 2, "name": "Ada", "start": "08:30", "score": 42, "ratio": 1.5}
    ]
    assert "rowid" not in response.json["columns"]


def test_rowid_follows_the_row_through_sorting(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string={"_shape": "objects", "_sort_desc": "name"})

    assert [row["rowid"] for row in response.json["rows"]] == [3, 2]


def test_visibility_toggles_drop_fields(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    hidden = client.get(endpoint, query_string={"_shape": "objects", "_rowid": "hide"})
    totalless = client.get(endpoint, query_string={"_total": "hide"})

    assert all("rowid" not in row for row in hidden.json["rows"])
    assert "total" not in totalless.json


@pytest.mark.parametrize(
    "params",
    [
        {"_size": 0},
        {"_size": -1},
        {"_size": "many"},
        {"_offset": -1},
        {"_offset": "1.5"},
        {"_shape": "tuples"},
        {"_rowid": "show"},
        {"_total": ""},
        {"_sort": ""},
        {"_sort": "nope"},
        {"_sort_desc": "nope"},
    ],
)
def test_control_parameter_errors(client, params):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string=params)

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]


@pytest.mark.parametrize(
    "name", ["_size", "_offset", "_shape", "_sort", "_sort_desc", "_rowid", "_total"]
)
def test_repeated_control_parameters_are_rejected(client, name):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert client.get(f"{endpoint}?{name}=1&{name}=1").status_code == 400


def test_explicit_charset_decodes_payload(client):
    client.payloads["https://example.com/cp1252.csv"] = "city,temp\nMünchen,21\n".encode("cp1252")

    response = convert(client, source="https://example.com/cp1252.csv", charset="cp1252")

    assert client.get(response.json["endpoint"]).json["rows"] == [["München", 21]]


@pytest.mark.parametrize(
    ("params", "status"),
    [
        ({}, 400),
        ({"source": "not-a-url"}, 400),
        ({"source": "https://example.com/data.csv", "charset": "klingon-8"}, 400),
        ({"source": "https://example.com/missing.csv"}, 404),
    ],
)
def test_convert_errors(client, params, status):
    response = convert(client, **params)

    assert response.status_code == status
    assert response.json["ok"] is False
    assert response.json["error"]


def test_non_tabular_source_is_rejected(client):
    client.payloads["https://example.com/page.html"] = b"<html>\n<body>hello</body>\n</html>\n"

    assert convert(client, source="https://example.com/page.html").status_code == 400


def test_single_row_source_is_rejected(client):
    client.payloads["https://example.com/header.csv"] = b"name,score\n"

    assert convert(client, source="https://example.com/header.csv").status_code == 400


def test_unknown_dataset_and_route_return_json_404(client):
    for path in ("/datasets/deadbeef", "/nope"):
        response = client.get(path)
        assert response.status_code == 404
        assert response.json["ok"] is False


def test_responses_carry_cors_headers(client):
    assert client.get("/nope").headers["Access-Control-Allow-Origin"] == "*"


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_delimiters_are_detected(delimiter):
    text = delimiter.join(["a", "b", "c"]) + "\n" + delimiter.join(["1", "2", "3"])

    assert detect_delimiter(text) == delimiter


@pytest.mark.parametrize(
    ("field", "expected"),
    [("42", 42), ("-3", -3), ("1.5", 1.5), ("2e3", 2000.0), ("08:30", "08:30"), ("n/a", "n/a")],
)
def test_value_inference(field, expected):
    assert infer_value(field) == expected


def test_encoding_detection_falls_back_to_latin1():
    assert decode("café".encode("utf-8")) == "café"
    assert decode("café".encode("utf-8-sig")) == "café"
    assert decode("café".encode("latin-1")) == "café"


def test_transport_failure_reports_missing_source(monkeypatch):
    def explode(url, timeout):
        raise requests.ConnectionError("no route to host")

    monkeypatch.setattr("fetcher.requests.get", explode)
    client = create_app().test_client()

    assert convert(client, source="https://example.invalid/x.csv").status_code == 404


def staff_dataset(client):
    """Convert a mixed text/number source and return its dataset endpoint."""
    client.payloads["https://example.com/staff.csv"] = (
        b"name,role,score\n"
        b"Ada,engineer,42\n"
        b"ada,Engineer,7\n"
        b"Grace,engineering lead,n/a\n"
        b"Kay,analyst,7.5\n"
    )
    return convert(client, source="https://example.com/staff.csv").json["endpoint"]


@pytest.mark.parametrize(
    ("params", "expected"),
    [
        ({"name__exact": "Ada"}, ["Ada"]),
        ({"name__exact": "ada"}, ["ada"]),
        ({"name__exact": "AD"}, []),
        ({"role__contains": "engineer"}, ["Ada", "Grace"]),
        ({"role__contains": "Engineer"}, ["ada"]),
        ({"score__less": "42"}, ["ada", "Kay"]),
        ({"score__greater": "7"}, ["Ada", "Kay"]),
        ({"score__exact": "42"}, ["Ada"]),
    ],
)
def test_comparators_select_rows(client, params, expected):
    endpoint = staff_dataset(client)

    response = client.get(endpoint, query_string=params)

    assert [row[0] for row in response.json["rows"]] == expected


def test_numeric_comparators_skip_non_numeric_values(client):
    endpoint = staff_dataset(client)

    for params in ({"score__less": "1000"}, {"score__greater": "-1000"}):
        response = client.get(endpoint, query_string=params)
        assert [row[0] for row in response.json["rows"]] == ["Ada", "ada", "Kay"]


def test_filters_are_combined_with_and(client):
    endpoint = staff_dataset(client)

    response = client.get(
        endpoint, query_string={"role__contains": "engineer", "score__greater": "10"}
    )

    assert [row[0] for row in response.json["rows"]] == ["Ada"]


def test_filtering_precedes_sorting_and_pagination(client):
    endpoint = staff_dataset(client)

    response = client.get(
        endpoint, query_string={"score__less": "42", "_sort_desc": "score", "_size": 1}
    )

    assert [row[0] for row in response.json["rows"]] == ["Kay"]


def test_total_counts_filtered_rows_before_pagination(client):
    endpoint = staff_dataset(client)

    response = client.get(
        endpoint, query_string={"role__contains": "engineer", "_size": 1}
    )

    assert len(response.json["rows"]) == 1
    assert response.json["total"] == 2


def test_plain_and_control_parameters_are_not_filters(client):
    endpoint = staff_dataset(client)

    response = client.get(f"{endpoint}?name=nobody&role=ghost&__exact=nobody")

    assert response.status_code == 200
    assert len(response.json["rows"]) == 4


@pytest.mark.parametrize(
    "query",
    [
        "name__matches=Ada",
        "name__=Ada",
        "score__less=soon",
        "score__greater=soon",
        "nickname__exact=Ada",
        "NAME__exact=Ada",
        "name__exact=Ada&name__exact=Grace",
    ],
)
def test_filter_errors(client, query):
    endpoint = staff_dataset(client)

    response = client.get(f"{endpoint}?{query}")

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]


def test_query_timeout_is_a_bad_request(client, monkeypatch):
    endpoint = staff_dataset(client)
    monkeypatch.setattr("filters.QUERY_BUDGET_SECONDS", -1.0)

    response = client.get(endpoint, query_string={"name__exact": "Ada"})

    assert response.status_code == 400
    assert response.json["ok"] is False


def export(client, endpoint, **params):
    return client.get(f"{endpoint}/export", query_string=params)


def csv_lines(response):
    """Read an export response back as a list of rows of text."""
    return list(csv.reader(io.StringIO(response.get_data(as_text=True))))


def test_export_serves_a_csv_attachment(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = export(client, endpoint)

    assert response.status_code == 200
    assert response.headers["Content-Type"] == "text/csv"
    dataset_id = endpoint.rsplit("/", 1)[-1]
    assert response.headers["Content-Disposition"] == f'attachment; filename="{dataset_id}.csv"'


def test_export_keeps_source_columns_and_values(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert csv_lines(export(client, endpoint)) == [
        ["name", "start", "score", "ratio"],
        ["Ada", "08:30", "42", "1.5"],
        ["Grace", "9:15", "7", "0.25"],
    ]


def test_export_applies_filters_then_sorting_then_pagination(client):
    endpoint = staff_dataset(client)

    response = export(client, endpoint, score__less="42", _sort_desc="score", _size=1)

    assert csv_lines(response) == [["name", "role", "score"], ["Kay", "analyst", "7.5"]]


def test_export_ignores_the_presentation_controls(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    shaped = export(client, endpoint, _shape="objects", _rowid="hide", _total="hide")

    assert csv_lines(shaped) == csv_lines(export(client, endpoint))


def test_export_quotes_values_holding_the_delimiter(client):
    client.payloads["https://example.com/quoted.csv"] = b'name,note\nAda,"one, two"\n'
    endpoint = convert(client, source="https://example.com/quoted.csv").json["endpoint"]

    assert csv_lines(export(client, endpoint))[1] == ["Ada", "one, two"]


def test_export_errors_use_the_json_envelope(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    unknown = export(client, "/datasets/deadbeef")
    invalid = export(client, endpoint, _size="many")

    assert (unknown.status_code, invalid.status_code) == (404, 400)
    assert unknown.json["ok"] is False and invalid.json["ok"] is False


def upload(client, payload, field="file", filename="data.csv", **params):
    return client.post(
        "/upload",
        query_string=params,
        data={field: (io.BytesIO(payload), filename)},
        content_type="multipart/form-data",
    )


@pytest.mark.parametrize("field", ["file", "attachment"])
def test_upload_publishes_a_queryable_dataset(client, field):
    response = upload(client, SAMPLE_CSV, field=field)

    assert response.status_code == 200
    assert response.json["ok"] is True
    rows = client.get(response.json["endpoint"]).json["rows"]
    assert rows == [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]


def test_upload_identifies_a_dataset_by_its_bytes(client):
    first = upload(client, SAMPLE_CSV, filename="one.csv")
    same = upload(client, SAMPLE_CSV, field="attachment", filename="two.csv")
    other = upload(client, b"city,temp\nOslo,3\n")

    assert first.json["endpoint"] == same.json["endpoint"]
    assert other.json["endpoint"] != first.json["endpoint"]


def test_upload_honours_the_charset_parameter(client):
    payload = "city,temp\nMünchen,21\n".encode("cp1252")

    response = upload(client, payload, charset="cp1252")

    assert client.get(response.json["endpoint"]).json["rows"] == [["München", 21]]


def test_non_multipart_upload_is_unsupported_media(client):
    response = client.post("/upload", data=SAMPLE_CSV, content_type="text/csv")

    assert response.status_code == 415
    assert response.json["ok"] is False


def test_upload_without_a_file_field_is_a_bad_request(client):
    response = client.post(
        "/upload", data={"payload": (io.BytesIO(SAMPLE_CSV), "data.csv")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 400
    assert response.json["ok"] is False


def test_malformed_multipart_body_is_a_bad_request(client):
    response = client.post(
        "/upload",
        data=b'--edge\r\nContent-Disposition: form-data; name="file"\r\n\r\ntruncated',
        content_type="multipart/form-data; boundary=edge",
    )

    assert response.status_code == 400
    assert response.json["ok"] is False


def test_unreadable_upload_is_a_bad_request(client):
    assert upload(client, b"<html>\n<body>hello</body>\n</html>\n").status_code == 400


def xlsx_bytes(*sheets: list[list]) -> bytes:
    """Serialise the given worksheets, first one first, into an .xlsx workbook."""
    workbook = openpyxl.Workbook()
    for index, rows in enumerate(sheets):
        sheet = workbook.worksheets[0] if index == 0 else workbook.create_sheet()
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def xls_bytes(*sheets: list[list]) -> bytes:
    """Serialise the given worksheets, first one first, into an .xls workbook."""
    workbook = xlwt.Workbook()
    for index, rows in enumerate(sheets):
        sheet = workbook.add_sheet(f"sheet{index}")
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                sheet.write(row_index, column_index, value)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


STAFF_SHEET = [["name", "start", "score"], ["Ada", "08:30", 42], ["Grace", "9:15", 7.5]]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_workbooks_convert_like_csv(client, workbook):
    client.payloads["https://example.com/staff.xlsx"] = workbook(STAFF_SHEET)
    endpoint = convert(client, source="https://example.com/staff.xlsx").json["endpoint"]

    response = client.get(endpoint)

    assert response.json["columns"] == ["name", "start", "score"]
    assert response.json["rows"] == [["Ada", "08:30", 42], ["Grace", "9:15", 7.5]]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_workbooks_upload_and_export(client, workbook):
    endpoint = upload(client, workbook(STAFF_SHEET), filename="staff.xls").json["endpoint"]

    assert csv_lines(export(client, endpoint)) == [
        ["name", "start", "score"],
        ["Ada", "08:30", "42"],
        ["Grace", "9:15", "7.5"],
    ]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_only_the_first_worksheet_is_ingested(client, workbook):
    payload = workbook(STAFF_SHEET, [["city", "temp"], ["Oslo", 3]])

    endpoint = upload(client, payload, filename="staff.xls").json["endpoint"]

    assert client.get(endpoint).json["columns"] == ["name", "start", "score"]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_workbook_rows_keep_their_sheet_line_numbers(client, workbook):
    endpoint = upload(client, workbook(STAFF_SHEET), filename="staff.xls").json["endpoint"]

    response = client.get(endpoint, query_string={"_shape": "objects"})

    assert [row["rowid"] for row in response.json["rows"]] == [2, 3]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_a_first_sheet_without_data_rows_is_rejected(client, workbook):
    payload = workbook([["name", "score"]], [["city", "temp"], ["Oslo", 3]])

    assert upload(client, payload, filename="header.xls").status_code == 400


def test_workbook_charset_is_not_validated(client):
    payload = xlsx_bytes(STAFF_SHEET)

    assert upload(client, payload, filename="staff.xlsx", charset="klingon-8").status_code == 200


def test_unrecognised_archive_is_a_bad_request(client):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")

    response = upload(client, stream.getvalue(), filename="bundle.zip")

    assert response.status_code == 400
    assert response.json["ok"] is False


REVISED_CSV = b"name,start,score,ratio\nKay,10:00,99,2.5\n"


def rows_of(client, endpoint):
    return client.get(endpoint).json["rows"]


def configured_client(client, monkeypatch, **settings):
    """A second client over the same fetch stub, under the given settings."""
    for name, value in settings.items():
        monkeypatch.setenv(name, value)
    fresh = create_app().test_client()
    fresh.payloads, fresh.fetched = client.payloads, client.fetched
    return fresh


def test_repeated_conversion_is_served_from_the_cache(client):
    first = convert(client, source="https://example.com/data.csv")
    client.payloads["https://example.com/data.csv"] = REVISED_CSV
    second = convert(client, source="https://example.com/data.csv")

    assert second.json == first.json
    assert client.fetched == ["https://example.com/data.csv"]
    assert rows_of(client, second.json["endpoint"]) == [
        ["Ada", "08:30", 42, 1.5],
        ["Grace", "9:15", 7, 0.25],
    ]


def test_force_replaces_the_cached_dataset(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]
    client.payloads["https://example.com/data.csv"] = REVISED_CSV

    forced = client.get("/convert?source=https://example.com/data.csv&force")

    assert forced.json == {"ok": True, "endpoint": endpoint}
    assert len(client.fetched) == 2
    assert rows_of(client, endpoint) == [["Kay", "10:00", 99, 2.5]]


@pytest.mark.parametrize("flag", ["force=", "force=1", "force=true", "force&force"])
def test_force_with_a_value_or_repeated_is_rejected(client, flag):
    response = client.get(f"/convert?source=https://example.com/data.csv&{flag}")

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]
    assert client.fetched == []


def test_disabled_caching_reingests_every_request(client, monkeypatch):
    fresh = configured_client(client, monkeypatch, CACHE_ENABLED="off")
    endpoint = convert(fresh, source="https://example.com/data.csv").json["endpoint"]
    fresh.payloads["https://example.com/data.csv"] = REVISED_CSV

    assert convert(fresh, source="https://example.com/data.csv").json["endpoint"] == endpoint
    assert fresh.fetched == ["https://example.com/data.csv"] * 2
    assert rows_of(fresh, endpoint) == [["Kay", "10:00", 99, 2.5]]


def test_force_changes_nothing_when_caching_is_disabled(client, monkeypatch):
    fresh = configured_client(client, monkeypatch, CACHE_ENABLED="off")
    endpoint = convert(fresh, source="https://example.com/data.csv").json["endpoint"]

    forced = fresh.get("/convert?source=https://example.com/data.csv&force")

    assert forced.json == {"ok": True, "endpoint": endpoint}
    assert fresh.fetched == ["https://example.com/data.csv"] * 2


@pytest.mark.parametrize(
    ("setting", "cached"),
    [("1", True), ("TRUE", True), ("Yes", True), ("on", True),
     ("0", False), ("false", False), ("NO", False), ("Off", False)],
)
def test_cache_setting_words_are_case_insensitive(client, monkeypatch, setting, cached):
    fresh = configured_client(client, monkeypatch, CACHE_ENABLED=setting)
    convert(fresh, source="https://example.com/data.csv")
    convert(fresh, source="https://example.com/data.csv")

    assert len(fresh.fetched) == (1 if cached else 2)


@pytest.mark.parametrize("setting", ["", "2", "maybe", "enabled"])
def test_invalid_cache_setting_fails_startup(monkeypatch, setting):
    monkeypatch.setenv("CACHE_ENABLED", setting)

    with pytest.raises(ValueError):
        create_app()


@pytest.mark.parametrize(
    ("payload", "status"),
    [(None, 404), (b"<html>\n<body>hello</body>\n</html>\n", 400)],
)
def test_failed_reingestion_keeps_the_prior_dataset(client, payload, status):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]
    if payload is None:
        del client.payloads["https://example.com/data.csv"]
    else:
        client.payloads["https://example.com/data.csv"] = payload

    failed = client.get("/convert?source=https://example.com/data.csv&force")

    assert failed.status_code == status
    assert failed.json["ok"] is False
    assert rows_of(client, endpoint) == [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]


SAMPLE_ROWS = [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]


def write_config(tmp_path, body: str) -> str:
    """Write a DATAGATE_CONFIG file and return its path."""
    path = tmp_path / "datagate.conf"
    path.write_text(body, encoding="utf-8")
    return str(path)


def test_defaults_apply_when_nothing_is_configured():
    settings = config.load({})

    assert settings.max_source_size is None
    assert settings.origin_allowlist == ()
    assert settings.require_tls is False
    assert settings.cache_enabled is True
    assert settings.storage_dir == config.DEFAULT_STORAGE_DIR


def test_config_file_is_read_around_blank_lines_and_comments(tmp_path):
    path = write_config(
        tmp_path,
        "# datagate settings\n"
        "\n"
        "MAX_SOURCE_SIZE = 2048\n"
        "ORIGIN_ALLOWLIST=example.com, eu.example.org\n"
        "REQUIRE_TLS=Yes\n"
        "CACHE_ENABLED=off\n"
        "STORAGE_DIR=/var/lib/datagate\n",
    )

    settings = config.load({"DATAGATE_CONFIG": path})

    assert settings.max_source_size == 2048
    assert settings.origin_allowlist == ("example.com", "eu.example.org")
    assert settings.require_tls is True
    assert settings.cache_enabled is False
    assert settings.storage_dir == pathlib.Path("/var/lib/datagate")


def test_environment_variables_outrank_the_config_file(tmp_path):
    path = write_config(tmp_path, "MAX_SOURCE_SIZE=10\nREQUIRE_TLS=false\n")

    settings = config.load(
        {"DATAGATE_CONFIG": path, "MAX_SOURCE_SIZE": "99", "REQUIRE_TLS": "on"}
    )

    assert settings.max_source_size == 99
    assert settings.require_tls is True


@pytest.mark.parametrize("word", [" on ", "TRUE", "Yes", "1"])
def test_boolean_words_are_trimmed_and_case_insensitive(word):
    assert config.load({"REQUIRE_TLS": word}).require_tls is True


@pytest.mark.parametrize(
    "environ",
    [
        {"MAX_SOURCE_SIZE": "plenty"},
        {"MAX_SOURCE_SIZE": "-1"},
        {"MAX_SOURCE_SIZE": "1.5"},
        {"REQUIRE_TLS": "maybe"},
    ],
)
def test_invalid_settings_are_rejected(environ):
    with pytest.raises(ConfigError):
        config.load(environ)


@pytest.mark.parametrize("line", ["MAX_SOURCE_SIZE", "MAX_SIZE=10", "hello there"])
def test_invalid_config_file_lines_are_rejected(tmp_path, line):
    path = write_config(tmp_path, f"REQUIRE_TLS=true\n{line}\n")

    with pytest.raises(ConfigError):
        config.load({"DATAGATE_CONFIG": path})


def test_unreadable_config_file_is_rejected(tmp_path):
    with pytest.raises(ConfigError):
        config.load({"DATAGATE_CONFIG": str(tmp_path / "absent.conf")})


def test_invalid_config_fails_startup(monkeypatch):
    monkeypatch.setenv("MAX_SOURCE_SIZE", "plenty")

    with pytest.raises(ConfigError):
        create_app()


def test_config_file_settings_reach_the_server(client, monkeypatch, tmp_path):
    monkeypatch.setenv("DATAGATE_CONFIG", write_config(tmp_path, "REQUIRE_TLS=true\n"))
    fresh = configured_client(client, monkeypatch)

    endpoint = convert(fresh, source="https://example.com/data.csv").json["endpoint"]

    assert endpoint.startswith("https://localhost/datasets/")


def limited_client(client, monkeypatch, limit: int):
    return configured_client(client, monkeypatch, MAX_SOURCE_SIZE=str(limit))


def test_a_source_of_exactly_the_limit_is_accepted(client, monkeypatch):
    fresh = limited_client(client, monkeypatch, len(SAMPLE_CSV))

    assert convert(fresh, source="https://example.com/data.csv").status_code == 200
    assert upload(fresh, SAMPLE_CSV).status_code == 200


def test_a_source_over_the_limit_is_rejected(client, monkeypatch):
    fresh = limited_client(client, monkeypatch, len(SAMPLE_CSV) - 1)

    response = convert(fresh, source="https://example.com/data.csv")

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]


def test_an_upload_over_the_limit_is_rejected(client, monkeypatch):
    fresh = limited_client(client, monkeypatch, len(SAMPLE_CSV) - 1)

    response = upload(fresh, SAMPLE_CSV)

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]


def test_an_oversized_source_leaves_no_dataset_behind(client, monkeypatch):
    fresh = limited_client(client, monkeypatch, 1)
    convert(fresh, source="https://example.com/data.csv")

    dataset_id = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert fresh.get(dataset_id).status_code == 404


def test_an_unset_limit_accepts_any_size(client):
    endpoint = big_dataset(client)

    assert client.get(endpoint).json["total"] == 250


def guarded_client(client, monkeypatch, allowlist="example.com,internal.test"):
    return configured_client(client, monkeypatch, ORIGIN_ALLOWLIST=allowlist)


@pytest.mark.parametrize(
    "referer",
    [
        "https://example.com/dashboard",
        "https://EXAMPLE.COM/dashboard",
        "https://eu.example.com/",
        "http://internal.test:8080/page",
    ],
)
def test_allowed_referers_are_routed(client, monkeypatch, referer):
    fresh = guarded_client(client, monkeypatch)

    response = fresh.get(
        "/convert",
        query_string={"source": "https://example.com/data.csv"},
        headers={"Referer": referer},
    )

    assert response.status_code == 200


@pytest.mark.parametrize(
    "referer",
    [
        "https://notexample.com/dashboard",
        "https://example.com.attacker.net/",
        "https://example.org/",
        "https://internal.test.co/",
        "nonsense",
    ],
)
def test_referers_outside_the_allowlist_are_forbidden(client, monkeypatch, referer):
    fresh = guarded_client(client, monkeypatch)

    response = fresh.get(
        "/convert",
        query_string={"source": "https://example.com/data.csv"},
        headers={"Referer": referer},
    )

    assert response.status_code == 403
    assert response.json["ok"] is False
    assert response.json["error"]
    assert fresh.fetched == []


def test_a_missing_referer_is_forbidden(client, monkeypatch):
    fresh = guarded_client(client, monkeypatch)

    response = convert(fresh, source="https://example.com/data.csv")

    assert response.status_code == 403
    assert response.json["ok"] is False
    assert response.json["error"]


def test_the_allowlist_is_checked_before_routing(client, monkeypatch):
    fresh = guarded_client(client, monkeypatch)

    assert fresh.get("/nowhere").status_code == 403
    assert client.get("/nowhere").status_code == 404


def test_uploads_and_dataset_routes_are_guarded_too(client, monkeypatch):
    fresh = guarded_client(client, monkeypatch)
    allowed = {"Referer": "https://example.com/"}
    endpoint = fresh.post(
        "/upload",
        data={"file": (io.BytesIO(SAMPLE_CSV), "data.csv")},
        content_type="multipart/form-data",
        headers=allowed,
    ).json["endpoint"]

    assert fresh.get(endpoint).status_code == 403
    assert fresh.get(endpoint, headers=allowed).status_code == 200


def test_without_an_allowlist_a_request_needs_no_referer(client):
    assert convert(client, source="https://example.com/data.csv").status_code == 200


def test_endpoints_are_relative_without_tls(client):
    converted = convert(client, source="https://example.com/data.csv")
    uploaded = upload(client, SAMPLE_CSV)

    assert converted.json["endpoint"].startswith("/datasets/")
    assert uploaded.json["endpoint"].startswith("/datasets/")


def test_endpoints_are_absolute_https_urls_with_tls(client, monkeypatch):
    path = convert(client, source="https://example.com/data.csv").json["endpoint"]
    fresh = configured_client(client, monkeypatch, REQUIRE_TLS="true")

    converted = convert(fresh, source="https://example.com/data.csv")
    uploaded = upload(fresh, SAMPLE_CSV)

    assert converted.json["endpoint"] == f"https://localhost{path}"
    assert uploaded.json["endpoint"].startswith("https://localhost/datasets/")


def test_the_storage_directory_is_created_when_missing(client, monkeypatch, tmp_path):
    directory = tmp_path / "missing" / "datasets"

    configured_client(client, monkeypatch, STORAGE_DIR=str(directory))

    assert directory.is_dir()


def test_datasets_outlive_the_server_that_converted_them(client, monkeypatch, tmp_path):
    directory = str(tmp_path / "persisted")
    first = configured_client(client, monkeypatch, STORAGE_DIR=directory)
    endpoint = convert(first, source="https://example.com/data.csv").json["endpoint"]

    restarted = configured_client(client, monkeypatch, STORAGE_DIR=directory)
    response = restarted.get(endpoint)

    assert response.json["rows"] == SAMPLE_ROWS
    assert convert(restarted, source="https://example.com/data.csv").json["endpoint"] == endpoint
    assert restarted.fetched == ["https://example.com/data.csv"]


def test_uploaded_datasets_are_persisted_too(client, monkeypatch, tmp_path):
    directory = str(tmp_path / "persisted")
    first = configured_client(client, monkeypatch, STORAGE_DIR=directory)
    endpoint = upload(first, SAMPLE_CSV).json["endpoint"]

    restarted = configured_client(client, monkeypatch, STORAGE_DIR=directory)

    assert restarted.get(endpoint).json["rows"] == SAMPLE_ROWS


def test_a_separate_storage_directory_starts_empty(client, monkeypatch, tmp_path):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    elsewhere = configured_client(client, monkeypatch, STORAGE_DIR=str(tmp_path / "other"))

    assert elsewhere.get(endpoint).status_code == 404


def test_an_empty_config_file_setting_names_no_file():
    assert config.load({"DATAGATE_CONFIG": ""}).require_tls is False


ENRICH_CSV = b"name,score,ratio,mixed\nAda,42,1.5,3\nGrace,,0.25,2.5\nAda,7,1.5,\n"
ENRICH_SOURCE = "https://example.com/enrich.csv"


def enrich(client, source=ENRICH_SOURCE, **params):
    """Convert a source with enrichment on and return its dataset body."""
    endpoint = convert(client, source=source, enrich="yes", **params).json["endpoint"]
    return client.get(endpoint).json


@pytest.fixture
def enrichable(client):
    """A client whose fetch table holds the profiling sample."""
    client.payloads[ENRICH_SOURCE] = ENRICH_CSV
    return client


def test_enriched_conversion_summarises_the_dataset(enrichable):
    body = enrich(enrichable)

    assert body["dataset_summary"] == {"filetype": "csv", "row_count": 3, "column_count": 4}


def test_enriched_conversion_profiles_every_column(enrichable):
    details = enrich(enrichable)["column_details"]

    assert details == {
        "name": {"type": "text", "distinct_count": 2, "missing_count": 0},
        "score": {"type": "integer", "distinct_count": 2, "missing_count": 1},
        "ratio": {"type": "float", "distinct_count": 2, "missing_count": 0},
        "mixed": {"type": "number", "distinct_count": 2, "missing_count": 1},
    }


def test_enrichment_leaves_the_conversion_response_alone(enrichable):
    plain = convert(enrichable, source="https://example.com/data.csv")
    enriched = convert(enrichable, source=ENRICH_SOURCE, enrich="yes")

    assert enriched.status_code == 200
    assert set(enriched.json) == set(plain.json) == {"ok", "endpoint"}
    assert enriched.json["ok"] is True
    assert enriched.json["endpoint"].startswith("/datasets/")


def test_enrichment_leaves_the_rows_and_the_query_fields_alone(client):
    plain = client.get(convert(client, source="https://example.com/data.csv").json["endpoint"])
    enriched = enrich(client, source="https://example.com/data.csv")

    assert enriched["ok"] is True
    assert enriched["columns"] == plain.json["columns"] == ["name", "start", "score", "ratio"]
    assert enriched["rows"] == plain.json["rows"] == SAMPLE_ROWS
    assert enriched["total"] == plain.json["total"] == 2
    assert enriched["query_ms"] >= 0


def test_a_plain_conversion_carries_no_metadata(enrichable):
    endpoint = convert(enrichable, source=ENRICH_SOURCE).json["endpoint"]

    body = enrichable.get(endpoint).json

    assert "dataset_summary" not in body
    assert "column_details" not in body


@pytest.mark.parametrize(
    "trigger",
    ["enrich=YES", "enrich=Yes", "enrich=no", "enrich=1", "enrich=true", "enrich=", "enrich",
     "enrich=yes&enrich=yes", "enrich=yes&enrich=no"],
)
def test_only_an_exact_single_enrich_yes_enriches(enrichable, trigger):
    response = enrichable.get(f"/convert?source={ENRICH_SOURCE}&{trigger}")

    assert response.status_code == 200
    assert "dataset_summary" not in enrichable.get(response.json["endpoint"]).json


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_enriched_workbooks_are_summarised_but_not_profiled(enrichable, workbook):
    enrichable.payloads["https://example.com/staff.xlsx"] = workbook(STAFF_SHEET)

    body = enrich(enrichable, source="https://example.com/staff.xlsx")

    assert body["dataset_summary"] == {"filetype": "excel", "row_count": 2, "column_count": 3}
    assert "column_details" not in body


def test_enriching_a_cached_dataset_reingests_and_upgrades_it(enrichable):
    convert(enrichable, source=ENRICH_SOURCE)

    body = enrich(enrichable)

    assert enrichable.fetched == [ENRICH_SOURCE] * 2
    assert body["dataset_summary"]["filetype"] == "csv"
    assert body["column_details"]["name"]["type"] == "text"


def test_an_enriched_dataset_is_served_from_the_cache(enrichable):
    first = enrich(enrichable)
    again = enrich(enrichable)

    assert enrichable.fetched == [ENRICH_SOURCE]
    assert again["dataset_summary"] == first["dataset_summary"]
    assert again["column_details"] == first["column_details"]


def test_a_plain_request_keeps_the_stored_enrichment(enrichable):
    endpoint = convert(enrichable, source=ENRICH_SOURCE, enrich="yes").json["endpoint"]

    convert(enrichable, source=ENRICH_SOURCE)

    assert enrichable.fetched == [ENRICH_SOURCE]
    assert "dataset_summary" in enrichable.get(endpoint).json


def test_reingestion_recomputes_the_metadata(enrichable):
    enrich(enrichable)
    enrichable.payloads[ENRICH_SOURCE] = REVISED_CSV

    forced = enrichable.get(f"/convert?source={ENRICH_SOURCE}&enrich=yes&force")
    body = enrichable.get(forced.json["endpoint"]).json

    assert body["dataset_summary"] == {"filetype": "csv", "row_count": 1, "column_count": 4}
    assert body["column_details"]["start"] == {
        "type": "text",
        "distinct_count": 1,
        "missing_count": 0,
    }


def test_a_forced_plain_conversion_drops_the_stored_enrichment(enrichable):
    endpoint = convert(enrichable, source=ENRICH_SOURCE, enrich="yes").json["endpoint"]

    enrichable.get(f"/convert?source={ENRICH_SOURCE}&force")

    assert "dataset_summary" not in enrichable.get(endpoint).json


def test_disabled_caching_still_follows_the_enrich_parameter(enrichable, monkeypatch):
    fresh = configured_client(enrichable, monkeypatch, CACHE_ENABLED="off")
    endpoint = convert(fresh, source=ENRICH_SOURCE, enrich="yes").json["endpoint"]

    assert "dataset_summary" in fresh.get(endpoint).json

    convert(fresh, source=ENRICH_SOURCE)

    assert "dataset_summary" not in fresh.get(endpoint).json


def test_enrichment_is_a_convert_parameter_only(enrichable):
    uploaded = upload(enrichable, ENRICH_CSV, enrich="yes").json["endpoint"]
    endpoint = convert(enrichable, source=ENRICH_SOURCE).json["endpoint"]

    assert "dataset_summary" not in enrichable.get(uploaded).json
    assert "dataset_summary" not in enrichable.get(endpoint, query_string={"enrich": "yes"}).json


@pytest.mark.parametrize(
    ("payload", "status"),
    [(None, 404), (b"<html>\n<body>hello</body>\n</html>\n", 400)],
)
def test_a_failed_enriched_reingestion_keeps_the_stored_metadata(enrichable, payload, status):
    endpoint = convert(enrichable, source=ENRICH_SOURCE, enrich="yes").json["endpoint"]
    if payload is None:
        del enrichable.payloads[ENRICH_SOURCE]
    else:
        enrichable.payloads[ENRICH_SOURCE] = payload

    failed = enrichable.get(f"/convert?source={ENRICH_SOURCE}&enrich=yes&force")

    assert failed.status_code == status
    assert failed.json["ok"] is False
    assert enrichable.get(endpoint).json["dataset_summary"]["row_count"] == 3


def test_enrichment_outlives_the_server_that_computed_it(enrichable, monkeypatch, tmp_path):
    directory = str(tmp_path / "persisted")
    first = configured_client(enrichable, monkeypatch, STORAGE_DIR=directory)
    endpoint = convert(first, source=ENRICH_SOURCE, enrich="yes").json["endpoint"]

    restarted = configured_client(enrichable, monkeypatch, STORAGE_DIR=directory)

    assert "column_details" in restarted.get(endpoint).json
    assert convert(restarted, source=ENRICH_SOURCE, enrich="yes").status_code == 200
    assert restarted.fetched == [ENRICH_SOURCE]


def test_a_fully_blank_column_reads_as_text_with_nothing_distinct(client):
    client.payloads["https://example.com/blank.csv"] = b"name,note\nAda,\nGrace,\n"

    details = enrich(client, source="https://example.com/blank.csv")["column_details"]

    assert details["note"] == {"type": "text", "distinct_count": 0, "missing_count": 2}
