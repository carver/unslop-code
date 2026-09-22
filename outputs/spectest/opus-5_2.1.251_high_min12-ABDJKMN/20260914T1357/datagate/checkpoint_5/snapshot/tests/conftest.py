"""Shared fixtures: a controllable origin HTTP server + a live datagate server."""
import http.server
import itertools
import json
import os
import socket
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATAGATE = os.path.join(ROOT, "datagate.py")


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


def wait_for_port(host, port, timeout=25.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            with socket.create_connection((host, port), timeout=0.5):
                return True
        except OSError:
            time.sleep(0.05)
    return False


def port_is_free(host, port):
    s = socket.socket()
    s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    try:
        s.bind((host, port))
        return True
    except OSError:
        return False
    finally:
        s.close()


# --------------------------------------------------------------------------
# Origin server: serves fixture payloads (bytes) with chosen status/content-type
# --------------------------------------------------------------------------
class Origin:
    def __init__(self):
        self.routes = {}
        self.hits = {}
        self.port = None

    @property
    def base(self):
        return "http://127.0.0.1:%d" % self.port

    def add(self, path, body, status=200, content_type="text/csv"):
        """Register a payload; returns its absolute URL."""
        if isinstance(body, str):
            body = body.encode("utf-8")
        if not path.startswith("/"):
            path = "/" + path
        self.routes[path] = (status, content_type, body)
        return self.base + path

    def add_text(self, path, text, charset="utf-8", content_type=None):
        if content_type is None:
            content_type = "text/csv"
        return self.add(path, text.encode(charset), content_type=content_type)


class _Handler(http.server.BaseHTTPRequestHandler):
    origin = None
    protocol_version = "HTTP/1.1"

    def do_GET(self):  # noqa: N802
        path = self.path.split("?")[0]
        o = self.origin
        o.hits[path] = o.hits.get(path, 0) + 1
        if path not in o.routes:
            body = b"not found"
            self.send_response(404)
            self.send_header("Content-Type", "text/plain")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        status, ctype, body = o.routes[path]
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_HEAD(self):  # noqa: N802
        self.do_GET()

    def log_message(self, *args):
        pass


@pytest.fixture(scope="session")
def origin():
    o = Origin()
    handler = type("H", (_Handler,), {"origin": o})
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    o.port = srv.server_address[1]
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    yield o
    srv.shutdown()
    srv.server_close()


# --------------------------------------------------------------------------
# datagate server under test
# --------------------------------------------------------------------------
class Client:
    def __init__(self, base):
        self.base = base

    def raw(self, path, method="GET", headers=None, timeout=30):
        url = path if path.startswith("http") else self.base + path
        req = urllib.request.Request(url, method=method, headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.status, dict(r.headers), r.read()
        except urllib.error.HTTPError as e:
            return e.code, dict(e.headers), e.read()

    def get(self, path, **kw):
        """Returns (status, headers, parsed_json_or_None)."""
        status, headers, body = self.raw(path, **kw)
        try:
            data = json.loads(body.decode("utf-8"))
        except Exception:
            data = None
        return status, headers, data

    def json(self, path, **kw):
        return self.get(path, **kw)[2]

    def convert(self, source=None, charset=None, extra=""):
        from urllib.parse import quote
        q = []
        if source is not None:
            q.append("source=" + quote(source, safe=""))
        if charset is not None:
            q.append("charset=" + quote(charset, safe=""))
        qs = "&".join(q)
        if extra:
            qs = (qs + "&" + extra) if qs else extra
        return self.get("/convert" + ("?" + qs if qs else ""))


def _spawn(port, address="127.0.0.1", args=None, tag="server"):
    log = open(os.path.join(ROOT, "tests", "_%s.log" % tag), "ab")
    if args is None:
        args = ["start", "--port", str(port), "--address", address]
    proc = subprocess.Popen(
        [sys.executable, DATAGATE] + args,
        cwd=ROOT, stdout=log, stderr=log, stdin=subprocess.DEVNULL,
    )
    return proc


def spawn_env(port, env_overrides=None, address="127.0.0.1", tag="env"):
    """Start datagate with extra environment variables (e.g. CACHE_ENABLED)."""
    env = dict(os.environ)
    for key, value in (env_overrides or {}).items():
        if value is None:
            env.pop(key, None)
        else:
            env[key] = value
    log = open(os.path.join(ROOT, "tests", "_%s.log" % tag), "ab")
    return subprocess.Popen(
        [sys.executable, DATAGATE, "start", "--port", str(port),
         "--address", address],
        cwd=ROOT, stdout=log, stderr=log, stdin=subprocess.DEVNULL, env=env,
    )


def stop(proc):
    """Terminate a spawned server by PID and reap it."""
    if proc.poll() is None:
        proc.terminate()
    try:
        proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait(timeout=10)


@pytest.fixture(scope="session")
def server():
    port = free_port()
    proc = _spawn(port)
    assert wait_for_port("127.0.0.1", port), "datagate did not start"
    yield Client("http://127.0.0.1:%d" % port)
    proc.terminate()
    try:
        proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        proc.kill()


# Globally unique so no two tests share a source URL: /convert caches by
# default, and a reused URL would serve an earlier test's dataset.
_PATH_COUNTER = itertools.count(1)


@pytest.fixture
def csv_url(origin):
    """Register a CSV payload under a unique path; return its URL."""
    def make(body, name=None, status=200, content_type="text/csv"):
        path = "/%s-%d.csv" % (name or "fixture", next(_PATH_COUNTER))
        return origin.add(path, body, status=status, content_type=content_type)

    return make


@pytest.fixture
def dataset(server, csv_url):
    """Convert a CSV body and return the parsed /datasets/<id> payload."""
    def make(body, charset=None, name=None, content_type="text/csv", query=""):
        url = csv_url(body, name=name, content_type=content_type)
        status, _, data = server.convert(url, charset=charset)
        assert status == 200, data
        endpoint = data["endpoint"]
        s2, _, payload = server.get(endpoint + query)
        assert s2 == 200, payload
        return payload
    return make


# --------------------------------------------------------------------------
# Multipart helpers for POST /upload
# --------------------------------------------------------------------------
BOUNDARY = "----datagateTestBoundary9d7f"


def multipart_body(parts, boundary=BOUNDARY):
    """Encode `parts` — (field, filename, bytes) or (field, value) — as bytes.

    A filename of None makes a plain (non-file) form field.
    """
    out = b""
    for part in parts:
        if len(part) == 2:
            field, filename, payload = part[0], None, part[1]
        else:
            field, filename, payload = part
        if isinstance(payload, str):
            payload = payload.encode("utf-8")
        head = 'Content-Disposition: form-data; name="%s"' % field
        if filename is not None:
            head += '; filename="%s"' % filename
            head += "\r\nContent-Type: application/octet-stream"
        out += ("--%s\r\n%s\r\n\r\n" % (boundary, head)).encode("utf-8")
        out += payload + b"\r\n"
    out += ("--%s--\r\n" % boundary).encode("utf-8")
    return out


class UploadClient(Client):
    def post(self, path, body, content_type=None, timeout=30):
        """POST raw bytes; returns (status, headers, parsed_json_or_None)."""
        url = path if path.startswith("http") else self.base + path
        headers = {}
        if content_type is not None:
            headers["Content-Type"] = content_type
        req = urllib.request.Request(url, data=body or b"", method="POST",
                                     headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                status, hdrs, raw = r.status, dict(r.headers), r.read()
        except urllib.error.HTTPError as e:
            status, hdrs, raw = e.code, dict(e.headers), e.read()
        try:
            data = json.loads(raw.decode("utf-8"))
        except Exception:
            data = None
        return status, hdrs, data

    def upload(self, payload, field="file", filename="data.csv", extra=None,
               boundary=BOUNDARY, content_type=None, path="/upload"):
        """Upload one file part (plus optional extra parts) to /upload."""
        parts = []
        if field is not None and payload is not None:
            parts.append((field, filename, payload))
        parts.extend(extra or [])
        body = multipart_body(parts, boundary=boundary)
        if content_type is None:
            content_type = "multipart/form-data; boundary=%s" % boundary
        return self.post(path, body, content_type=content_type)


@pytest.fixture
def up(server):
    """The session server, wrapped with POST/upload helpers."""
    return UploadClient(server.base)


# --------------------------------------------------------------------------
# Spreadsheet fixture builders
# --------------------------------------------------------------------------
def make_xlsx(sheets):
    """Build .xlsx bytes from [(title, [[cell, ...], ...]), ...]."""
    import io as _io

    import openpyxl

    wb = openpyxl.Workbook()
    for index, (title, rows) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        for row in rows:
            ws.append(list(row))
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_xls(sheets):
    """Build legacy .xls bytes from [(title, [[cell, ...], ...]), ...]."""
    import io as _io

    import xlwt

    wb = xlwt.Workbook()
    for title, rows in sheets:
        ws = wb.add_sheet(title)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value is None:
                    continue
                ws.write(r, c, value)
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_csv_bytes(raw):
    """Parse exported CSV bytes into a list of rows of strings."""
    import csv as _csv
    import io as _io

    text = raw.decode("utf-8")
    return [row for row in _csv.reader(_io.StringIO(text, newline=""))
            if row != []]
