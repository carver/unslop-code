"""End-to-end tests for the datagate HTTP API and CSV pipeline."""

import csv
import io
import zipfile

import openpyxl
import pytest
import requests
import xlwt

from app import create_app
from errors import ApiError
from tabular import decode, detect_delimiter, infer_value

SAMPLE_CSV = b"name,start,score,ratio\nAda,08:30,42,1.5\nGrace,9:15,7,0.25\n"


@pytest.fixture
def client(monkeypatch):
    """A test client whose fetches are served from an in-test payload table."""
    payloads = {"https://example.com/data.csv": SAMPLE_CSV}
    fetched = []

    def fake_fetch(url):
        fetched.append(url)
        if url not in payloads:
            raise ApiError(f"Source unreachable: {url}", 404)
        return payloads[url]

    monkeypatch.setattr("app.fetch", fake_fetch)
    app = create_app()
    app.config["TESTING"] = True
    client = app.test_client()
    client.payloads = payloads
    client.fetched = fetched
    return client


def convert(client, **params):
    return client.get("/convert", query_string=params)


def big_dataset(client):
    """Convert a 250 row source and return its dataset endpoint."""
    rows = "\n".join(f"row{index};{index}" for index in range(250))
    client.payloads["https://example.com/big.csv"] = f"label;value\n{rows}\n".encode()
    return convert(client, source="https://example.com/big.csv").json["endpoint"]


def test_convert_returns_stable_endpoint(client):
    first = convert(client, source="https://example.com/data.csv")
    second = convert(client, source="https://example.com/data.csv")

    assert first.status_code == 200
    assert first.json["ok"] is True
    assert first.json["endpoint"] == second.json["endpoint"]


def test_dataset_preserves_order_and_infers_types(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint)

    assert response.status_code == 200
    assert response.json["columns"] == ["name", "start", "score", "ratio"]
    assert response.json["rows"] == [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]
    assert response.json["total"] == 2
    assert response.json["query_ms"] >= 0


def test_rows_are_capped_at_one_hundred_by_default(client):
    endpoint = big_dataset(client)

    assert len(client.get(endpoint).json["rows"]) == 100
    assert len(client.get(endpoint, query_string={"_size": 5}).json["rows"]) == 5


def test_size_beyond_the_row_count_returns_everything(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert len(client.get(endpoint, query_string={"_size": 500}).json["rows"]) == 2


def test_offset_skips_rows_and_total_ignores_pagination(client):
    endpoint = big_dataset(client)

    response = client.get(endpoint, query_string={"_offset": 248, "_size": 10})

    assert response.json["rows"] == [["row248", 248], ["row249", 249]]
    assert response.json["total"] == 250


def test_offset_past_the_end_returns_no_rows(client):
    endpoint = big_dataset(client)

    assert client.get(endpoint, query_string={"_offset": 400}).json["rows"] == []


def test_total_counts_rows_before_pagination(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert client.get(endpoint, query_string={"_size": 1}).json["total"] == 2


def test_sorting_runs_before_pagination(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    ascending = client.get(endpoint, query_string={"_sort": "score", "_size": 1})
    descending = client.get(endpoint, query_string={"_sort_desc": "score", "_size": 1})

    assert ascending.json["rows"] == [["Grace", "9:15", 7, 0.25]]
    assert descending.json["rows"] == [["Ada", "08:30", 42, 1.5]]


def test_descending_sort_wins_over_ascending(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string={"_sort": "name", "_sort_desc": "name"})

    assert [row[0] for row in response.json["rows"]] == ["Grace", "Ada"]


def test_sorting_keeps_tied_rows_in_source_order(client):
    client.payloads["https://example.com/ties.csv"] = (
        b"name,team\nAda,blue\nGrace,blue\nKay,amber\n"
    )
    endpoint = convert(client, source="https://example.com/ties.csv").json["endpoint"]

    for params in ({"_sort": "team"}, {"_sort_desc": "team"}):
        response = client.get(endpoint, query_string=params)
        assert [row[0] for row in response.json["rows"] if row[1] == "blue"] == ["Ada", "Grace"]


def test_object_shape_carries_rowid(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string={"_shape": "objects", "_size": 1})

    assert response.json["rows"] == [
        {"rowid": 2, "name": "Ada", "start": "08:30", "score": 42, "ratio": 1.5}
    ]
    assert "rowid" not in response.json["columns"]


def test_rowid_follows_the_row_through_sorting(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string={"_shape": "objects", "_sort_desc": "name"})

    assert [row["rowid"] for row in response.json["rows"]] == [3, 2]


def test_visibility_toggles_drop_fields(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    hidden = client.get(endpoint, query_string={"_shape": "objects", "_rowid": "hide"})
    totalless = client.get(endpoint, query_string={"_total": "hide"})

    assert all("rowid" not in row for row in hidden.json["rows"])
    assert "total" not in totalless.json


@pytest.mark.parametrize(
    "params",
    [
        {"_size": 0},
        {"_size": -1},
        {"_size": "many"},
        {"_offset": -1},
        {"_offset": "1.5"},
        {"_shape": "tuples"},
        {"_rowid": "show"},
        {"_total": ""},
        {"_sort": ""},
        {"_sort": "nope"},
        {"_sort_desc": "nope"},
    ],
)
def test_control_parameter_errors(client, params):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = client.get(endpoint, query_string=params)

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]


@pytest.mark.parametrize(
    "name", ["_size", "_offset", "_shape", "_sort", "_sort_desc", "_rowid", "_total"]
)
def test_repeated_control_parameters_are_rejected(client, name):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert client.get(f"{endpoint}?{name}=1&{name}=1").status_code == 400


def test_explicit_charset_decodes_payload(client):
    client.payloads["https://example.com/cp1252.csv"] = "city,temp\nMünchen,21\n".encode("cp1252")

    response = convert(client, source="https://example.com/cp1252.csv", charset="cp1252")

    assert client.get(response.json["endpoint"]).json["rows"] == [["München", 21]]


@pytest.mark.parametrize(
    ("params", "status"),
    [
        ({}, 400),
        ({"source": "not-a-url"}, 400),
        ({"source": "https://example.com/data.csv", "charset": "klingon-8"}, 400),
        ({"source": "https://example.com/missing.csv"}, 404),
    ],
)
def test_convert_errors(client, params, status):
    response = convert(client, **params)

    assert response.status_code == status
    assert response.json["ok"] is False
    assert response.json["error"]


def test_non_tabular_source_is_rejected(client):
    client.payloads["https://example.com/page.html"] = b"<html>\n<body>hello</body>\n</html>\n"

    assert convert(client, source="https://example.com/page.html").status_code == 400


def test_single_row_source_is_rejected(client):
    client.payloads["https://example.com/header.csv"] = b"name,score\n"

    assert convert(client, source="https://example.com/header.csv").status_code == 400


def test_unknown_dataset_and_route_return_json_404(client):
    for path in ("/datasets/deadbeef", "/nope"):
        response = client.get(path)
        assert response.status_code == 404
        assert response.json["ok"] is False


def test_responses_carry_cors_headers(client):
    assert client.get("/nope").headers["Access-Control-Allow-Origin"] == "*"


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_delimiters_are_detected(delimiter):
    text = delimiter.join(["a", "b", "c"]) + "\n" + delimiter.join(["1", "2", "3"])

    assert detect_delimiter(text) == delimiter


@pytest.mark.parametrize(
    ("field", "expected"),
    [("42", 42), ("-3", -3), ("1.5", 1.5), ("2e3", 2000.0), ("08:30", "08:30"), ("n/a", "n/a")],
)
def test_value_inference(field, expected):
    assert infer_value(field) == expected


def test_encoding_detection_falls_back_to_latin1():
    assert decode("café".encode("utf-8")) == "café"
    assert decode("café".encode("utf-8-sig")) == "café"
    assert decode("café".encode("latin-1")) == "café"


def test_transport_failure_reports_missing_source(monkeypatch):
    def explode(url, timeout):
        raise requests.ConnectionError("no route to host")

    monkeypatch.setattr("fetcher.requests.get", explode)
    client = create_app().test_client()

    assert convert(client, source="https://example.invalid/x.csv").status_code == 404


def staff_dataset(client):
    """Convert a mixed text/number source and return its dataset endpoint."""
    client.payloads["https://example.com/staff.csv"] = (
        b"name,role,score\n"
        b"Ada,engineer,42\n"
        b"ada,Engineer,7\n"
        b"Grace,engineering lead,n/a\n"
        b"Kay,analyst,7.5\n"
    )
    return convert(client, source="https://example.com/staff.csv").json["endpoint"]


@pytest.mark.parametrize(
    ("params", "expected"),
    [
        ({"name__exact": "Ada"}, ["Ada"]),
        ({"name__exact": "ada"}, ["ada"]),
        ({"name__exact": "AD"}, []),
        ({"role__contains": "engineer"}, ["Ada", "Grace"]),
        ({"role__contains": "Engineer"}, ["ada"]),
        ({"score__less": "42"}, ["ada", "Kay"]),
        ({"score__greater": "7"}, ["Ada", "Kay"]),
        ({"score__exact": "42"}, ["Ada"]),
    ],
)
def test_comparators_select_rows(client, params, expected):
    endpoint = staff_dataset(client)

    response = client.get(endpoint, query_string=params)

    assert [row[0] for row in response.json["rows"]] == expected


def test_numeric_comparators_skip_non_numeric_values(client):
    endpoint = staff_dataset(client)

    for params in ({"score__less": "1000"}, {"score__greater": "-1000"}):
        response = client.get(endpoint, query_string=params)
        assert [row[0] for row in response.json["rows"]] == ["Ada", "ada", "Kay"]


def test_filters_are_combined_with_and(client):
    endpoint = staff_dataset(client)

    response = client.get(
        endpoint, query_string={"role__contains": "engineer", "score__greater": "10"}
    )

    assert [row[0] for row in response.json["rows"]] == ["Ada"]


def test_filtering_precedes_sorting_and_pagination(client):
    endpoint = staff_dataset(client)

    response = client.get(
        endpoint, query_string={"score__less": "42", "_sort_desc": "score", "_size": 1}
    )

    assert [row[0] for row in response.json["rows"]] == ["Kay"]


def test_total_counts_filtered_rows_before_pagination(client):
    endpoint = staff_dataset(client)

    response = client.get(
        endpoint, query_string={"role__contains": "engineer", "_size": 1}
    )

    assert len(response.json["rows"]) == 1
    assert response.json["total"] == 2


def test_plain_and_control_parameters_are_not_filters(client):
    endpoint = staff_dataset(client)

    response = client.get(f"{endpoint}?name=nobody&role=ghost&__exact=nobody")

    assert response.status_code == 200
    assert len(response.json["rows"]) == 4


@pytest.mark.parametrize(
    "query",
    [
        "name__matches=Ada",
        "name__=Ada",
        "score__less=soon",
        "score__greater=soon",
        "nickname__exact=Ada",
        "NAME__exact=Ada",
        "name__exact=Ada&name__exact=Grace",
    ],
)
def test_filter_errors(client, query):
    endpoint = staff_dataset(client)

    response = client.get(f"{endpoint}?{query}")

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]


def test_query_timeout_is_a_bad_request(client, monkeypatch):
    endpoint = staff_dataset(client)
    monkeypatch.setattr("filters.QUERY_BUDGET_SECONDS", -1.0)

    response = client.get(endpoint, query_string={"name__exact": "Ada"})

    assert response.status_code == 400
    assert response.json["ok"] is False


def export(client, endpoint, **params):
    return client.get(f"{endpoint}/export", query_string=params)


def csv_lines(response):
    """Read an export response back as a list of rows of text."""
    return list(csv.reader(io.StringIO(response.get_data(as_text=True))))


def test_export_serves_a_csv_attachment(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    response = export(client, endpoint)

    assert response.status_code == 200
    assert response.headers["Content-Type"] == "text/csv"
    dataset_id = endpoint.rsplit("/", 1)[-1]
    assert response.headers["Content-Disposition"] == f'attachment; filename="{dataset_id}.csv"'


def test_export_keeps_source_columns_and_values(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    assert csv_lines(export(client, endpoint)) == [
        ["name", "start", "score", "ratio"],
        ["Ada", "08:30", "42", "1.5"],
        ["Grace", "9:15", "7", "0.25"],
    ]


def test_export_applies_filters_then_sorting_then_pagination(client):
    endpoint = staff_dataset(client)

    response = export(client, endpoint, score__less="42", _sort_desc="score", _size=1)

    assert csv_lines(response) == [["name", "role", "score"], ["Kay", "analyst", "7.5"]]


def test_export_ignores_the_presentation_controls(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    shaped = export(client, endpoint, _shape="objects", _rowid="hide", _total="hide")

    assert csv_lines(shaped) == csv_lines(export(client, endpoint))


def test_export_quotes_values_holding_the_delimiter(client):
    client.payloads["https://example.com/quoted.csv"] = b'name,note\nAda,"one, two"\n'
    endpoint = convert(client, source="https://example.com/quoted.csv").json["endpoint"]

    assert csv_lines(export(client, endpoint))[1] == ["Ada", "one, two"]


def test_export_errors_use_the_json_envelope(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]

    unknown = export(client, "/datasets/deadbeef")
    invalid = export(client, endpoint, _size="many")

    assert (unknown.status_code, invalid.status_code) == (404, 400)
    assert unknown.json["ok"] is False and invalid.json["ok"] is False


def upload(client, payload, field="file", filename="data.csv", **params):
    return client.post(
        "/upload",
        query_string=params,
        data={field: (io.BytesIO(payload), filename)},
        content_type="multipart/form-data",
    )


@pytest.mark.parametrize("field", ["file", "attachment"])
def test_upload_publishes_a_queryable_dataset(client, field):
    response = upload(client, SAMPLE_CSV, field=field)

    assert response.status_code == 200
    assert response.json["ok"] is True
    rows = client.get(response.json["endpoint"]).json["rows"]
    assert rows == [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]


def test_upload_identifies_a_dataset_by_its_bytes(client):
    first = upload(client, SAMPLE_CSV, filename="one.csv")
    same = upload(client, SAMPLE_CSV, field="attachment", filename="two.csv")
    other = upload(client, b"city,temp\nOslo,3\n")

    assert first.json["endpoint"] == same.json["endpoint"]
    assert other.json["endpoint"] != first.json["endpoint"]


def test_upload_honours_the_charset_parameter(client):
    payload = "city,temp\nMünchen,21\n".encode("cp1252")

    response = upload(client, payload, charset="cp1252")

    assert client.get(response.json["endpoint"]).json["rows"] == [["München", 21]]


def test_non_multipart_upload_is_unsupported_media(client):
    response = client.post("/upload", data=SAMPLE_CSV, content_type="text/csv")

    assert response.status_code == 415
    assert response.json["ok"] is False


def test_upload_without_a_file_field_is_a_bad_request(client):
    response = client.post(
        "/upload", data={"payload": (io.BytesIO(SAMPLE_CSV), "data.csv")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 400
    assert response.json["ok"] is False


def test_malformed_multipart_body_is_a_bad_request(client):
    response = client.post(
        "/upload",
        data=b'--edge\r\nContent-Disposition: form-data; name="file"\r\n\r\ntruncated',
        content_type="multipart/form-data; boundary=edge",
    )

    assert response.status_code == 400
    assert response.json["ok"] is False


def test_unreadable_upload_is_a_bad_request(client):
    assert upload(client, b"<html>\n<body>hello</body>\n</html>\n").status_code == 400


def xlsx_bytes(*sheets: list[list]) -> bytes:
    """Serialise the given worksheets, first one first, into an .xlsx workbook."""
    workbook = openpyxl.Workbook()
    for index, rows in enumerate(sheets):
        sheet = workbook.worksheets[0] if index == 0 else workbook.create_sheet()
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def xls_bytes(*sheets: list[list]) -> bytes:
    """Serialise the given worksheets, first one first, into an .xls workbook."""
    workbook = xlwt.Workbook()
    for index, rows in enumerate(sheets):
        sheet = workbook.add_sheet(f"sheet{index}")
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                sheet.write(row_index, column_index, value)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


STAFF_SHEET = [["name", "start", "score"], ["Ada", "08:30", 42], ["Grace", "9:15", 7.5]]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_workbooks_convert_like_csv(client, workbook):
    client.payloads["https://example.com/staff.xlsx"] = workbook(STAFF_SHEET)
    endpoint = convert(client, source="https://example.com/staff.xlsx").json["endpoint"]

    response = client.get(endpoint)

    assert response.json["columns"] == ["name", "start", "score"]
    assert response.json["rows"] == [["Ada", "08:30", 42], ["Grace", "9:15", 7.5]]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_workbooks_upload_and_export(client, workbook):
    endpoint = upload(client, workbook(STAFF_SHEET), filename="staff.xls").json["endpoint"]

    assert csv_lines(export(client, endpoint)) == [
        ["name", "start", "score"],
        ["Ada", "08:30", "42"],
        ["Grace", "9:15", "7.5"],
    ]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_only_the_first_worksheet_is_ingested(client, workbook):
    payload = workbook(STAFF_SHEET, [["city", "temp"], ["Oslo", 3]])

    endpoint = upload(client, payload, filename="staff.xls").json["endpoint"]

    assert client.get(endpoint).json["columns"] == ["name", "start", "score"]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_workbook_rows_keep_their_sheet_line_numbers(client, workbook):
    endpoint = upload(client, workbook(STAFF_SHEET), filename="staff.xls").json["endpoint"]

    response = client.get(endpoint, query_string={"_shape": "objects"})

    assert [row["rowid"] for row in response.json["rows"]] == [2, 3]


@pytest.mark.parametrize("workbook", [xlsx_bytes, xls_bytes])
def test_a_first_sheet_without_data_rows_is_rejected(client, workbook):
    payload = workbook([["name", "score"]], [["city", "temp"], ["Oslo", 3]])

    assert upload(client, payload, filename="header.xls").status_code == 400


def test_workbook_charset_is_not_validated(client):
    payload = xlsx_bytes(STAFF_SHEET)

    assert upload(client, payload, filename="staff.xlsx", charset="klingon-8").status_code == 200


def test_unrecognised_archive_is_a_bad_request(client):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")

    response = upload(client, stream.getvalue(), filename="bundle.zip")

    assert response.status_code == 400
    assert response.json["ok"] is False


REVISED_CSV = b"name,start,score,ratio\nKay,10:00,99,2.5\n"


def rows_of(client, endpoint):
    return client.get(endpoint).json["rows"]


def uncached_client(client, monkeypatch, setting="off"):
    """A second client over the same fetch stub, with caching switched off."""
    monkeypatch.setenv("CACHE_ENABLED", setting)
    fresh = create_app().test_client()
    fresh.payloads, fresh.fetched = client.payloads, client.fetched
    return fresh


def test_repeated_conversion_is_served_from_the_cache(client):
    first = convert(client, source="https://example.com/data.csv")
    client.payloads["https://example.com/data.csv"] = REVISED_CSV
    second = convert(client, source="https://example.com/data.csv")

    assert second.json == first.json
    assert client.fetched == ["https://example.com/data.csv"]
    assert rows_of(client, second.json["endpoint"]) == [
        ["Ada", "08:30", 42, 1.5],
        ["Grace", "9:15", 7, 0.25],
    ]


def test_force_replaces_the_cached_dataset(client):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]
    client.payloads["https://example.com/data.csv"] = REVISED_CSV

    forced = client.get("/convert?source=https://example.com/data.csv&force")

    assert forced.json == {"ok": True, "endpoint": endpoint}
    assert len(client.fetched) == 2
    assert rows_of(client, endpoint) == [["Kay", "10:00", 99, 2.5]]


@pytest.mark.parametrize("flag", ["force=", "force=1", "force=true", "force&force"])
def test_force_with_a_value_or_repeated_is_rejected(client, flag):
    response = client.get(f"/convert?source=https://example.com/data.csv&{flag}")

    assert response.status_code == 400
    assert response.json["ok"] is False
    assert response.json["error"]
    assert client.fetched == []


def test_disabled_caching_reingests_every_request(client, monkeypatch):
    fresh = uncached_client(client, monkeypatch)
    endpoint = convert(fresh, source="https://example.com/data.csv").json["endpoint"]
    fresh.payloads["https://example.com/data.csv"] = REVISED_CSV

    assert convert(fresh, source="https://example.com/data.csv").json["endpoint"] == endpoint
    assert fresh.fetched == ["https://example.com/data.csv"] * 2
    assert rows_of(fresh, endpoint) == [["Kay", "10:00", 99, 2.5]]


def test_force_changes_nothing_when_caching_is_disabled(client, monkeypatch):
    fresh = uncached_client(client, monkeypatch)
    endpoint = convert(fresh, source="https://example.com/data.csv").json["endpoint"]

    forced = fresh.get("/convert?source=https://example.com/data.csv&force")

    assert forced.json == {"ok": True, "endpoint": endpoint}
    assert fresh.fetched == ["https://example.com/data.csv"] * 2


@pytest.mark.parametrize(
    ("setting", "cached"),
    [("1", True), ("TRUE", True), ("Yes", True), ("on", True),
     ("0", False), ("false", False), ("NO", False), ("Off", False)],
)
def test_cache_setting_words_are_case_insensitive(client, monkeypatch, setting, cached):
    fresh = uncached_client(client, monkeypatch, setting)
    convert(fresh, source="https://example.com/data.csv")
    convert(fresh, source="https://example.com/data.csv")

    assert len(fresh.fetched) == (1 if cached else 2)


@pytest.mark.parametrize("setting", ["", "2", "maybe", " true", "enabled"])
def test_invalid_cache_setting_fails_startup(monkeypatch, setting):
    monkeypatch.setenv("CACHE_ENABLED", setting)

    with pytest.raises(ValueError):
        create_app()


@pytest.mark.parametrize(
    ("payload", "status"),
    [(None, 404), (b"<html>\n<body>hello</body>\n</html>\n", 400)],
)
def test_failed_reingestion_keeps_the_prior_dataset(client, payload, status):
    endpoint = convert(client, source="https://example.com/data.csv").json["endpoint"]
    if payload is None:
        del client.payloads["https://example.com/data.csv"]
    else:
        client.payloads["https://example.com/data.csv"] = payload

    failed = client.get("/convert?source=https://example.com/data.csv&force")

    assert failed.status_code == status
    assert failed.json["ok"] is False
    assert rows_of(client, endpoint) == [["Ada", "08:30", 42, 1.5], ["Grace", "9:15", 7, 0.25]]
