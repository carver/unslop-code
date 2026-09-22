"""Builders for the in-memory ``.xlsx`` and ``.xls`` files the ingestion tests use."""

import io

import openpyxl
import xlwt

Grid = list[list]


def xlsx(*sheets: Grid) -> bytes:
    """Return an ``.xlsx`` workbook holding one worksheet per grid given, in order."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for index, grid in enumerate(sheets):
        worksheet = workbook.create_sheet(f"sheet{index}")
        for row in grid:
            worksheet.append(row)

    document = io.BytesIO()
    workbook.save(document)
    return document.getvalue()


def xls(*sheets: Grid) -> bytes:
    """Return an ``.xls`` workbook holding one worksheet per grid given, in order."""
    workbook = xlwt.Workbook()
    for index, grid in enumerate(sheets):
        worksheet = workbook.add_sheet(f"sheet{index}")
        for row_index, row in enumerate(grid):
            for column_index, cell in enumerate(row):
                worksheet.write(row_index, column_index, cell)

    document = io.BytesIO()
    workbook.save(document)
    return document.getvalue()
