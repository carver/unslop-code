"""Spec coverage for /datasets/<id>/export, /upload and multi-format ingestion."""
import csv, io, json, socket, subprocess, sys, threading, time
import http.server, socketserver, urllib.parse, urllib.request, urllib.error
import zipfile

import openpyxl, xlwt

CSV = ("name,age,score\n"
       "Alice,30,9.5\n"
       "bob,25,8\n"
       "Carol,25,7.25\n"
       "dave,40,\n"
       "Eve,25,9.5\n")
QUOTED = ('name,note,n\n'
          '"Smith, John","said ""hi""",3\n'
          '"Doe, Jane","line\nbreak",4\n')


def xlsx_bytes(sheets):
    wb = openpyxl.Workbook()
    first = True
    for title, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        first = False
        for row in rows:
            ws.append(row)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def xls_bytes(sheets):
    wb = xlwt.Workbook()
    for title, rows in sheets:
        ws = wb.add_sheet(title)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                ws.write(r, c, value)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


XLSX = xlsx_bytes([("first", [["city", "pop", "ratio"],
                              ["Paris", 2148000, 1.5],
                              ["Lyon", 513000, 0.9]]),
                   ("second", [["zzz"], ["other sheet"]])])
XLS = xls_bytes([("first", [["city", "pop", "ratio"],
                            ["Paris", 2148000, 1.5],
                            ["Lyon", 513000, 0.9]]),
                 ("second", [["zzz"], ["other sheet"]])])
XLSX_HEADER_ONLY = xlsx_bytes([("only", [["a", "b", "c"]])])
XLS_HEADER_ONLY = xls_bytes([("only", [["a", "b", "c"]])])
XLSX_EMPTY = xlsx_bytes([("only", [])])

_zbuf = io.BytesIO()
with zipfile.ZipFile(_zbuf, "w") as _z:
    _z.writestr("readme.txt", "just a zip, not a workbook")
PLAIN_ZIP = _zbuf.getvalue()
PDF = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n" + bytes(range(256)) * 4

FIXTURES = {
    "/t.csv": ("text/csv", CSV.encode()),
    "/book.xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", XLSX),
    "/book.xls": ("application/vnd.ms-excel", XLS),
    "/thin.xlsx": ("application/octet-stream", XLSX_HEADER_ONLY),
    "/thin.xls": ("application/octet-stream", XLS_HEADER_ONLY),
    "/plain.zip": ("application/zip", PLAIN_ZIP),
    "/doc.pdf": ("application/pdf", PDF),
    "/latin.csv": ("text/csv", "name;ville\nJos\xe9;M\xfcnch\xe9n\nRen\xe9;K\xf6ln\n".encode("latin-1")),
}


class H(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        p = urllib.parse.urlsplit(self.path).path
        ctype, body = FIXTURES.get(p, ("text/plain", b"nope"))
        code = 200 if p in FIXTURES else 404
        self.send_response(code); self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body))); self.end_headers()
        self.wfile.write(body)
    def log_message(self, *a): pass


def free_port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p

FP = free_port()
srv = socketserver.ThreadingTCPServer(("127.0.0.1", FP), H); srv.daemon_threads = True
threading.Thread(target=srv.serve_forever, daemon=True).start()
FIX = f"http://127.0.0.1:{FP}"

PORT = free_port()
proc = subprocess.Popen([sys.executable, "datagate.py", "start", "--port", str(PORT),
                         "--address", "127.0.0.1"], cwd="/workspace",
                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
BASE = f"http://127.0.0.1:{PORT}"
for _ in range(100):
    try: urllib.request.urlopen(BASE + "/", timeout=1); break
    except Exception: time.sleep(0.1)

fails = []

def check(name, cond, extra=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  -> {extra}"))
    if not cond: fails.append(name)

def raw_get(path):
    req = urllib.request.Request(BASE + path, headers={"Origin": "http://example.com"})
    try:
        r = urllib.request.urlopen(req, timeout=30)
        return r.status, r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, e.read(), dict(e.headers)

def get(path):
    s, body, h = raw_get(path)
    try: return s, json.loads(body.decode()), h
    except Exception: return s, {"_raw": body}, h

def post(path, body, ctype):
    req = urllib.request.Request(BASE + path, data=body, method="POST")
    if ctype is not None:
        req.add_header("Content-Type", ctype)
    try:
        r = urllib.request.urlopen(req, timeout=30)
        return r.status, json.loads(r.read().decode()), dict(r.headers)
    except urllib.error.HTTPError as e:
        raw = e.read()
        try: return e.code, json.loads(raw.decode()), dict(e.headers)
        except Exception: return e.code, {"_raw": raw}, dict(e.headers)

BOUNDARY = "----dgboundary9z"

def multipart(parts):
    """parts: list of (fieldname, filename|None, bytes)."""
    out = b""
    for name, filename, content in parts:
        out += f"--{BOUNDARY}\r\n".encode()
        disp = f'Content-Disposition: form-data; name="{name}"'
        if filename is not None:
            disp += f'; filename="{filename}"'
        out += disp.encode() + b"\r\n"
        out += b"Content-Type: application/octet-stream\r\n\r\n" if filename else b"\r\n"
        out += content + b"\r\n"
    out += f"--{BOUNDARY}--\r\n".encode()
    return out, f"multipart/form-data; boundary={BOUNDARY}"

def upload(parts, ctype=None, body=None):
    if body is None:
        body, auto = multipart(parts)
        ctype = ctype or auto
    return post("/upload", body, ctype)

def conv(src, charset=None):
    q = {"source": src}
    if charset is not None: q["charset"] = charset
    return get("/convert?" + urllib.parse.urlencode(q))

def rows_of(body):
    return list(csv.reader(io.StringIO(body.decode("utf-8"), newline="")))

# ---------------------------------------------------------------- export ----
_, b, _ = conv(FIX + "/t.csv")
EP = b["endpoint"]
DSID = EP.rsplit("/", 1)[1]

s, body, h = raw_get(EP + "/export")
check("export 200", s == 200, (s, body[:120]))
check("export content-type", h.get("Content-Type") == "text/csv", h.get("Content-Type"))
check("export content-disposition",
      h.get("Content-Disposition") == f'attachment; filename="{DSID}.csv"',
      h.get("Content-Disposition"))
check("export cors", h.get("Access-Control-Allow-Origin") == "*", h)
rows = rows_of(body)
check("export header is source order", rows[0] == ["name", "age", "score"], rows[0])
check("export all rows", rows[1:] == [["Alice", "30", "9.5"], ["bob", "25", "8"],
                                      ["Carol", "25", "7.25"], ["dave", "40", ""],
                                      ["Eve", "25", "9.5"]], rows[1:])

s, body, _ = raw_get(EP + "/export?_size=2")
check("export pagination", rows_of(body) == [["name", "age", "score"],
                                             ["Alice", "30", "9.5"],
                                             ["bob", "25", "8"]], rows_of(body))
s, body, _ = raw_get(EP + "/export?_size=2&_offset=2")
check("export offset", rows_of(body)[1:] == [["Carol", "25", "7.25"],
                                             ["dave", "40", ""]], rows_of(body))
s, body, _ = raw_get(EP + "/export?age__exact=25")
check("export filter", [r[0] for r in rows_of(body)[1:]] == ["bob", "Carol", "Eve"],
      rows_of(body))
s, body, _ = raw_get(EP + "/export?_sort=name")
check("export sort", [r[0] for r in rows_of(body)[1:]] == ["Alice", "Carol", "Eve", "bob", "dave"],
      rows_of(body))
s, body, _ = raw_get(EP + "/export?_sort_desc=age")
check("export sort desc", [r[1] for r in rows_of(body)[1:]] == ["40", "30", "25", "25", "25"],
      rows_of(body))
s, body, _ = raw_get(EP + "/export?age__exact=25&_sort_desc=score&_size=2")
check("export filter+sort+paginate",
      rows_of(body)[1:] == [["Eve", "25", "9.5"], ["bob", "25", "8"]], rows_of(body))

base = raw_get(EP + "/export")[1]
for ctl in ["_shape=objects", "_shape=lists", "_rowid=hide", "_total=hide",
            "_shape=objects&_rowid=hide&_total=hide"]:
    s2, body2, _ = raw_get(EP + "/export?" + ctl)
    check(f"export ignores {ctl}", s2 == 200 and body2 == base, body2[:120])

s, b, _ = get("/datasets/deadbeefdeadbeef/export")
check("export unknown 404", s == 404 and b.get("ok") is False and isinstance(b.get("error"), str), (s, b))
s, b, _ = get(EP + "/export?_size=0")
check("export bad control 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = get(EP + "/export?nope__exact=1")
check("export bad filter 400", s == 400 and b.get("ok") is False, (s, b))

# quoted round-trip
s, b, _ = upload([("file", "q.csv", QUOTED.encode())])
qep = b["endpoint"]
_, body, _ = raw_get(qep + "/export")
check("export quoting round-trips",
      rows_of(body) == [["name", "note", "n"], ["Smith, John", 'said "hi"', "3"],
                        ["Doe, Jane", "line\nbreak", "4"]], rows_of(body))

# ---------------------------------------------------------------- upload ----
s, b, h = upload([("file", "t.csv", CSV.encode())])
check("upload file field 200", s == 200 and b.get("ok") is True
      and b.get("endpoint", "").startswith("/datasets/"), (s, b))
uep = b["endpoint"]
s2, b2, _ = upload([("file", "t.csv", CSV.encode())])
check("upload deterministic id", b2.get("endpoint") == uep, (uep, b2))
s3, b3, _ = upload([("attachment", "other-name.csv", CSV.encode())])
check("upload attachment field 200", s3 == 200 and b3.get("ok") is True, (s3, b3))
check("same bytes same id (file vs attachment)", b3.get("endpoint") == uep, (uep, b3))
s4, b4, _ = upload([("file", "t.csv", (CSV + "Zoe,22,5\n").encode())])
check("different bytes different id", b4.get("endpoint") != uep, b4)

s, b, _ = get(uep)
check("uploaded dataset queryable", s == 200 and b["columns"] == ["name", "age", "score"]
      and b["rows"][0] == ["Alice", 30, 9.5] and b["total"] == 5, b)
_, body, h = raw_get(uep + "/export")
check("uploaded export", rows_of(body)[0] == ["name", "age", "score"], rows_of(body)[:1])

# non-multipart -> 415
s, b, _ = post("/upload", json.dumps({"file": "x"}).encode(), "application/json")
check("json upload 415", s == 415 and b.get("ok") is False and isinstance(b.get("error"), str), (s, b))
s, b, _ = post("/upload", CSV.encode(), "text/csv")
check("text/csv body 415", s == 415 and b.get("ok") is False, (s, b))
s, b, _ = post("/upload", urllib.parse.urlencode({"file": CSV}).encode(),
               "application/x-www-form-urlencoded")
check("urlencoded 415", s == 415 and b.get("ok") is False, (s, b))
s, b, _ = post("/upload", b"", None)
check("no content-type 415", s == 415, (s, b))

# malformed multipart / missing fields -> 400
body, ctype = multipart([("other", "t.csv", CSV.encode())])
s, b, _ = upload(None, ctype=ctype, body=body)
check("missing file field 400", s == 400 and b.get("ok") is False, (s, b))
body, ctype = multipart([("desc", None, b"hello")])
s, b, _ = upload(None, ctype=ctype, body=body)
check("no file part at all 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = upload(None, ctype="multipart/form-data", body=b"garbage without boundary")
check("multipart w/o boundary 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = upload(None, ctype=f"multipart/form-data; boundary={BOUNDARY}",
                 body=b"total nonsense not a part at all")
check("malformed multipart 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = upload(None, ctype=f"multipart/form-data; boundary={BOUNDARY}",
                 body=f"--{BOUNDARY}\r\nContent-Disposition: form-data; name=\"file\"".encode())
check("truncated multipart 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = upload([("file", "empty.csv", b"")])
check("empty file 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = upload([("file", "x.bin", PDF)])
check("unsupported format upload 400", s == 400 and b.get("ok") is False, (s, b))

# ---------------------------------------------------------- multi-format ----
for label, path in [("xlsx", "/book.xlsx"), ("xls", "/book.xls")]:
    s, b, _ = conv(FIX + path)
    check(f"convert {label} 200", s == 200 and b.get("ok") is True, (s, b))
    _, d, _ = get(b["endpoint"])
    check(f"convert {label} columns order", d["columns"] == ["city", "pop", "ratio"], d.get("columns"))
    check(f"convert {label} rows", d["rows"] == [["Paris", 2148000, 1.5], ["Lyon", 513000, 0.9]],
          d.get("rows"))
    check(f"convert {label} first sheet only", d["total"] == 2, d.get("total"))
    _, body, h = raw_get(b["endpoint"] + "/export")
    check(f"export {label}", rows_of(body) == [["city", "pop", "ratio"],
                                               ["Paris", "2148000", "1.5"],
                                               ["Lyon", "513000", "0.9"]], rows_of(body))

for label, blob in [("xlsx", XLSX), ("xls", XLS)]:
    s, b, _ = upload([("file", f"book.{label}", blob)])
    check(f"upload {label} 200", s == 200 and b.get("ok") is True, (s, b))
    _, d, _ = get(b["endpoint"])
    check(f"upload {label} parsed", d["columns"] == ["city", "pop", "ratio"]
          and d["rows"][0] == ["Paris", 2148000, 1.5], d)
    s2, b2, _ = upload([("attachment", "renamed.dat", blob)])
    check(f"upload {label} id from bytes", b2.get("endpoint") == b.get("endpoint"), (b, b2))

for label, path in [("xlsx", "/thin.xlsx"), ("xls", "/thin.xls")]:
    s, b, _ = conv(FIX + path)
    check(f"convert header-only {label} 400", s == 400 and b.get("ok") is False, (s, b))
for label, blob in [("xlsx", XLSX_HEADER_ONLY), ("xls", XLS_HEADER_ONLY), ("empty xlsx", XLSX_EMPTY)]:
    s, b, _ = upload([("file", "thin.x", blob)])
    check(f"upload header-only {label} 400", s == 400 and b.get("ok") is False, (s, b))

for label, path in [("plain zip", "/plain.zip"), ("pdf", "/doc.pdf")]:
    s, b, _ = conv(FIX + path)
    check(f"convert unsupported {label} 400", s == 400 and b.get("ok") is False, (s, b))
s, b, _ = upload([("file", "plain.zip", PLAIN_ZIP)])
check("upload unsupported zip 400", s == 400 and b.get("ok") is False, (s, b))

# charset applies to CSV only
s, b, _ = conv(FIX + "/latin.csv", charset="latin-1")
_, d, _ = get(b["endpoint"])
check("charset on csv", d["rows"][0] == ["José", "Münchén"], d.get("rows"))
s, b, _ = conv(FIX + "/book.xlsx", charset="latin-1")
check("charset ignored for xlsx", s == 200 and b.get("ok") is True, (s, b))
_, d, _ = get(b["endpoint"])
check("charset ignored for xlsx rows", d["rows"][0] == ["Paris", 2148000, 1.5], d.get("rows"))
s, b, _ = conv(FIX + "/book.xls", charset="utf-8")
check("charset ignored for xls", s == 200 and b.get("ok") is True, (s, b))
s, b, _ = upload([("file", "b.xlsx", XLSX)], ctype=None)
check("upload xlsx no charset needed", s == 200, (s, b))
s, b, _ = post("/upload?charset=latin-1", *reversed(list(reversed(multipart(
    [("file", "l.csv", "name;ville\nJos\xe9;K\xf6ln\n".encode("latin-1"))])))))
_, d, _ = get(b["endpoint"])
check("upload charset param", d["rows"][0] == ["José", "Köln"], d.get("rows"))
s, b, _ = conv(FIX + "/t.csv", charset="nonsense-9000")
check("bad charset still 400", s == 400, (s, b))

# text CSV served with a spreadsheet-ish name is still ingested as CSV
s, b, _ = upload([("file", "actually.xlsx", CSV.encode())])
check("csv bytes named .xlsx -> csv", s == 200 and b.get("ok") is True, (s, b))


# ------------------------------------------------------------- extra edges --
BIG = "i,sq\n" + "".join(f"{i},{i*i}\n" for i in range(250))
s, b, _ = upload([("file", "big.csv", BIG.encode())])
bep = b["endpoint"]
_, body, _ = raw_get(bep + "/export")
r = rows_of(body)
check("export default limit 100", len(r) == 101 and r[1] == ["0", "0"] and r[100] == ["99", "9801"],
      (len(r), r[:2]))
_, body, _ = raw_get(bep + "/export?_size=250")
check("export explicit big size", len(rows_of(body)) == 251, len(rows_of(body)))
_, body, _ = raw_get(bep + "/export?_size=3&_offset=248")
check("export tail window", rows_of(body)[1:] == [["248", "61504"], ["249", "62001"]],
      rows_of(body))

# non-ascii export is utf-8
s, b, _ = upload([("file", "u.csv", "name,ville\nJos\u00e9,M\u00fcnch\u00e9n\n".encode("utf-8"))])
_, body, h = raw_get(b["endpoint"] + "/export")
check("export utf-8 body", rows_of(body) == [["name", "ville"], ["José", "Münchén"]], body)

# HEAD on export
req = urllib.request.Request(BASE + EP + "/export", method="HEAD")
try:
    r = urllib.request.urlopen(req, timeout=10)
    hs, hh = r.status, dict(r.headers)
except urllib.error.HTTPError as e:
    hs, hh = e.code, dict(e.headers)
check("HEAD export 200 + headers", hs == 200 and hh.get("Content-Type") == "text/csv"
      and hh.get("Content-Disposition", "").startswith("attachment;"), (hs, hh))

# first sheet empty, later sheet tabular -> 400
mixed = xlsx_bytes([("blank", []), ("data", [["a", "b"], [1, 2]])])
s, b, _ = upload([("file", "mixed.xlsx", mixed)])
check("first sheet not tabular 400", s == 400 and b.get("ok") is False, (s, b))

# plain (non-file) multipart part named file still ingested
body, ctype = multipart([("file", None, CSV.encode())])
s, b, _ = upload(None, ctype=ctype, body=body)
check("plain form part named file", s == 200 and b.get("ok") is True, (s, b))

# spreadsheet with dates / bools / blanks
import datetime as _dt
rich = xlsx_bytes([("s", [["when", "flag", "note", "n"],
                          [_dt.date(2024, 5, 17), True, None, 0.25],
                          [_dt.datetime(2024, 5, 17, 8, 30), False, "x", 3]])])
s, b, _ = upload([("file", "rich.xlsx", rich)])
check("rich xlsx 200", s == 200, (s, b))
_, d, _ = get(b["endpoint"])
check("rich xlsx values", d["columns"] == ["when", "flag", "note", "n"]
      and d["rows"][0] == ["2024-05-17", True, "", 0.25]
      and d["rows"][1] == ["2024-05-17 08:30:00", False, "x", 3], d.get("rows"))
_, body, _ = raw_get(b["endpoint"] + "/export")
check("rich xlsx export", rows_of(body) == [["when", "flag", "note", "n"],
                                            ["2024-05-17", "true", "", "0.25"],
                                            ["2024-05-17 08:30:00", "false", "x", "3"]],
      rows_of(body))

# ragged sheet: short rows padded to the header width
ragged = xls_bytes([("s", [["a", "b", "c"], [1], [1, 2, 3]])])
s, b, _ = upload([("file", "ragged.xls", ragged)])
_, d, _ = get(b["endpoint"])
check("ragged xls padded", d["rows"] == [[1, "", ""], [1, 2, 3]], d.get("rows"))

# uploads and /convert share the dataset surface
s, b, _ = upload([("file", "t.csv", CSV.encode())])
s2, b2, _ = get(b["endpoint"] + "?_shape=objects&_size=1")
check("uploaded objects shape", s2 == 200 and b2["rows"][0] == {"rowid": 1, "name": "Alice",
                                                               "age": 30, "score": 9.5}, b2)

proc.terminate(); proc.wait(timeout=10); srv.shutdown()
print("\n%d failure(s): %s" % (len(fails), fails))
sys.exit(1 if fails else 0)
