"""Builders for the ``.xlsx`` and ``.xls`` payloads the format tests ingest."""

import io

import xlwt
from openpyxl import Workbook

Sheet = list[list]


def xlsx_bytes(*sheets: Sheet) -> bytes:
    """Pack each sheet of rows into an Office Open XML workbook."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, rows in enumerate(sheets):
        sheet = workbook.create_sheet(f"sheet{index}")
        for row in rows:
            sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xls_bytes(*sheets: Sheet) -> bytes:
    """Pack each sheet of rows into a legacy BIFF workbook."""
    workbook = xlwt.Workbook()
    for index, rows in enumerate(sheets):
        sheet = workbook.add_sheet(f"sheet{index}")
        for row_index, row in enumerate(rows):
            for column_index, cell in enumerate(row):
                sheet.write(row_index, column_index, cell)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
