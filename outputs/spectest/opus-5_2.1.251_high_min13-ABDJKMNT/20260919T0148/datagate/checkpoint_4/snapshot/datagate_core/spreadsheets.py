"""Reading the first worksheet of an `.xls` or `.xlsx` workbook into table rows."""

import contextlib
import io

import openpyxl
import xlrd

from .errors import DatagateError
from .formats import Format
from .tables import split_header

SOURCE = "the first worksheet"


def read_workbook(raw, workbook_format):
    """Return the header and rows of `raw`'s first worksheet, as text cells."""
    readers = {Format.XLSX: _read_xlsx, Format.XLS: _read_xls}
    return split_header(_trimmed(readers[workbook_format](raw)), SOURCE)


def _read_xlsx(raw):
    """Read the first worksheet of an `.xlsx` workbook as a grid of text cells.

    `data_only` asks openpyxl for the value cached alongside each formula rather
    than for the formula itself; an uncached formula reads as an empty cell.
    """
    with _rejecting_unreadable(".xlsx"):
        sheet = openpyxl.load_workbook(io.BytesIO(raw), data_only=True).worksheets[0]

    return [[cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]


def _read_xls(raw):
    """Read the first worksheet of a legacy `.xls` workbook as a grid of text cells."""
    with _rejecting_unreadable(".xls"):
        book = xlrd.open_workbook(file_contents=raw)
        sheet = book.sheet_by_index(0)

    return [
        [_xls_text(sheet.cell(row, column), book.datemode) for column in range(sheet.ncols)]
        for row in range(sheet.nrows)
    ]


def cell_text(value):
    """Render one cell as the text the CSV reader would have carried for it.

    Whole doubles lose their `.0` so a number reads the same whichever format holds
    it: `.xls` stores every number, integer or not, as a double (AMBIGUITIES T43).
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xls_text(cell, datemode):
    """Render an `.xls` cell, restoring the types xlrd flattens into doubles."""
    if cell.ctype == xlrd.XL_CELL_DATE:
        return str(xlrd.xldate_as_datetime(cell.value, datemode))
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return str(bool(cell.value))
    return cell_text(cell.value)


@contextlib.contextmanager
def _rejecting_unreadable(extension):
    """Turn any complaint from a workbook library into the spec's 400 for the format.

    The libraries raise their own hierarchies plus whatever the underlying zip or
    OLE2 reader raises, and a payload that merely shares a container signature with
    a workbook is an unrecognized format rather than a server fault.
    """
    try:
        yield
    except Exception as error:
        raise DatagateError(
            400, f"Unrecognized format: not a readable {extension} workbook."
        ) from error


def _trimmed(grid):
    """Drop the wholly empty trailing rows and columns of a worksheet's used range.

    Both formats report a used range that a cleared cell or stray formatting can
    inflate past the data, and that padding would otherwise surface as blank rows
    and unnamed columns (AMBIGUITIES T42).
    """
    rows = list(grid)
    while rows and not any(rows[-1]):
        rows.pop()

    width = max((_used_width(row) for row in rows), default=0)
    return [list(row[:width]) for row in rows]


def _used_width(row):
    """The cell count of `row` up to and including its last non-empty cell."""
    filled = [index for index, cell in enumerate(row) if cell]
    return filled[-1] + 1 if filled else 0
