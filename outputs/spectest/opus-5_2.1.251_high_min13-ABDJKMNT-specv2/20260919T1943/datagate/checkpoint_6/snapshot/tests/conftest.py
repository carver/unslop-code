"""Shared fixtures: a test client plus helpers for fake remote CSVs, workbooks and uploads."""

import io
import pathlib
import sys

import openpyxl
import pytest
import responses
import xlwt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from datagate_app.server import create_app  # noqa: E402

SOURCE_URL = "https://example.test/data.csv"

SIMPLE_CSV = "name,age\nada,36\ngrace,45\n"


@pytest.fixture(autouse=True)
def storage_dir(tmp_path, monkeypatch):
    """Give every app built during a test its own (not yet created) storage directory."""
    directory = tmp_path / "storage"
    monkeypatch.setenv("STORAGE_DIR", str(directory))
    return directory


@pytest.fixture
def configured(monkeypatch):
    """Build a client from an app that read the given settings at startup."""

    def _configured(**settings):
        for key, value in settings.items():
            monkeypatch.setenv(key, value)
        app = create_app()
        app.config.update(TESTING=True)
        return app.test_client()

    return _configured


@pytest.fixture
def client(configured):
    return configured()


@pytest.fixture
def remote():
    """Intercept outbound HTTP so tests never touch the network."""
    with responses.RequestsMock(assert_all_requests_are_fired=False) as mock:
        yield mock


@pytest.fixture
def serve(remote):
    """Publish a body at a URL and hand back that URL."""

    def _serve(body, url=SOURCE_URL, content_type="text/csv", status=200):
        remote.add(responses.GET, url, body=body, status=status, content_type=content_type)
        return url

    return _serve


@pytest.fixture
def dataset(client, serve):
    """Convert a CSV body and return the parsed `/datasets/<id>` payload."""

    def _dataset(body, url=SOURCE_URL, content_type="text/csv", query=""):
        endpoint = client.get("/convert", query_string={"source": serve(body, url, content_type)})
        assert endpoint.status_code == 200, endpoint.get_json()
        return client.get(endpoint.get_json()["endpoint"] + query).get_json()

    return _dataset


@pytest.fixture
def endpoint(client, serve):
    """Convert a CSV body and hand back its `/datasets/<id>` path."""

    def _endpoint(body=SIMPLE_CSV, url=SOURCE_URL, content_type="text/csv"):
        response = client.get("/convert", query_string={"source": serve(body, url, content_type)})
        assert response.status_code == 200, response.get_json()
        return response.get_json()["endpoint"]

    return _endpoint


@pytest.fixture
def query(client, endpoint):
    """Request `/datasets/<id>` with a control query string and return the response."""

    def _query(query_string, body=SIMPLE_CSV, url=SOURCE_URL):
        return client.get(endpoint(body, url) + query_string)

    return _query


def xlsx_bytes(sheets, active=0):
    """Build an `.xlsx` workbook from `{sheet_name: rows}`, first key first."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.active = active
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xls_bytes(sheets):
    """Build a legacy `.xls` workbook from `{sheet_name: rows}`, first key first."""
    workbook = xlwt.Workbook()
    for name, rows in sheets.items():
        sheet = workbook.add_sheet(name)
        for row_number, row in enumerate(rows):
            for column, cell in enumerate(row):
                if cell is not None:
                    sheet.write(row_number, column, cell)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


SIMPLE_SHEET = {"people": [["name", "age"], ["ada", 36], ["grace", 45]]}


@pytest.fixture
def upload(client):
    """POST a multipart upload and return the raw response."""

    def _upload(data, filename="data.csv", field="file", query=None, content_type=None):
        return client.post(
            "/upload",
            data={field: (io.BytesIO(data), filename)},
            content_type=content_type or "multipart/form-data",
            query_string=query or {},
        )

    return _upload


@pytest.fixture
def uploaded(upload, client):
    """Upload a payload and return the parsed `/datasets/<id>` payload."""

    def _uploaded(data, filename="data.csv", query=None, dataset_query=""):
        response = upload(data, filename=filename, query=query)
        assert response.status_code == 200, response.get_json()
        return client.get(response.get_json()["endpoint"] + dataset_query).get_json()

    return _uploaded
