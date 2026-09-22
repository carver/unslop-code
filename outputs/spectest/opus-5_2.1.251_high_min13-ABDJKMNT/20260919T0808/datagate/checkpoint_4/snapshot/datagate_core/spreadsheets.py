"""Reading the first worksheet of an `.xls` or `.xlsx` workbook as a grid of cells."""

import io
import zipfile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.compdoc import CompDocError

from datagate_core.errors import UnsupportedFormatError

#: Signatures the two workbook formats start with: a ZIP container, then OLE2 (see T35).
XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

#: What the readers raise when a container is damaged or holds something else entirely.
BROKEN_WORKBOOK = (zipfile.BadZipFile, InvalidFileException, KeyError, xlrd.XLRDError, CompDocError)


def is_workbook(data: bytes) -> bool:
    """True when the bytes carry a workbook signature rather than text."""
    return data.startswith((XLSX_MAGIC, XLS_MAGIC))


def read_first_sheet(data: bytes) -> list[list[str]]:
    """The first worksheet as rows of text cells, skipping rows that hold no content."""
    reader = _read_xlsx if data.startswith(XLSX_MAGIC) else _read_xls
    try:
        rows = reader(data)
    except BROKEN_WORKBOOK as error:
        raise UnsupportedFormatError(f"not a readable workbook: {error}")
    return [trimmed for trimmed in (_trim(row) for row in rows) if trimmed]


def _read_xlsx(data: bytes) -> list[list[str]]:
    """Read sheet one of an OOXML workbook, taking cached values over formulas."""
    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    return [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]


def _read_xls(data: bytes) -> list[list[str]]:
    """Read sheet one of a legacy BIFF workbook."""
    sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
    return [[_cell_text(cell) for cell in sheet.row_values(n)] for n in range(sheet.nrows)]


def _cell_text(value) -> str:
    """One cell as source text: blanks empty, whole numbers without a `.0` tail (T41)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _trim(row: list[str]) -> list[str]:
    """Drop the trailing blank cells a sheet's stored width pads its rows out with."""
    while row and not row[-1].strip():
        row = row[:-1]
    return row
