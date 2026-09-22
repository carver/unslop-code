#!/usr/bin/env python3
"""datagate -- turn remote CSV files into deterministic, queryable JSON datasets.

Usage:
    python datagate.py start --port <port> --address <address>
"""
from __future__ import annotations

import argparse
import codecs
import csv
import hashlib
import io
import json
import os
import re
import sys
import tempfile
import threading
import time
from urllib.parse import unquote_plus, urlparse

import requests
from flask import Flask, Request, Response, jsonify, request
from werkzeug.exceptions import HTTPException
from werkzeug.formparser import FormDataParser

DEFAULT_PORT = 8001
DEFAULT_ADDRESS = "127.0.0.1"

ROW_LIMIT = 100                 # spec: default row limit
DELIMITERS = (",", ";", "\t")   # spec: minimum supported delimiters
FETCH_TIMEOUT = 30              # seconds
MAX_BYTES = 64 * 1024 * 1024    # refuse absurdly large downloads
MAX_PROSE_TOKENS = 4            # single-column guard, see AMBIGUITIES T9

# Query controls on /datasets/<id>; repeating any of them is a 400.
CONTROL_PARAMS = ("_size", "_offset", "_shape", "_sort", "_sort_desc", "_rowid", "_total")
SHAPES = ("lists", "objects")

# Column filters: <column>__<comparator>=<value>.
COMPARATORS = ("exact", "contains", "less", "greater")
NUMERIC_COMPARATORS = ("less", "greater")
FILTER_SEP = "__"
ROWS_PER_DEADLINE_CHECK = 1024

# Configuration. Precedence: built-in defaults < DATAGATE_CONFIG file <
# direct environment variables. Booleans are strict, case-insensitive, and
# trimmed; anything else is a startup failure.
CONFIG_ENV = "DATAGATE_CONFIG"
CACHE_ENV = "CACHE_ENABLED"
BOOL_TRUE = ("1", "true", "yes", "on")
BOOL_FALSE = ("0", "false", "no", "off")
CACHE_TRUE = BOOL_TRUE
CACHE_FALSE = BOOL_FALSE
SETTING_NAMES = (
    "MAX_SOURCE_SIZE",
    "ORIGIN_ALLOWLIST",
    "REQUIRE_TLS",
    "STORAGE_DIR",
    CACHE_ENV,
)
COMMENT_PREFIX = "#"
LIST_SEPARATOR = ","
# `STORAGE_DIR` is documented as implementation-defined; a stable path outside
# the working tree keeps persisted datasets across restarts (AMBIGUITIES T62).
DEFAULT_STORAGE_DIR = os.path.join(tempfile.gettempdir(), "datagate-storage")
DATASET_FILE_SUFFIX = ".json"
DATASET_ID_RE = re.compile(r"^[0-9a-f]{1,64}$")
FORCE_PARAM = "force"

# Optional ingestion enrichment on /convert: one exact `enrich=yes`, no more.
ENRICH_PARAM = "enrich"
ENRICH_VALUE = "yes"
FILETYPE_CSV = "csv"
FILETYPE_EXCEL = "excel"
SPREADSHEET_FORMATS = ("xlsx", "xls")
TYPE_TEXT = "text"
TYPE_NUMBER = "number"
TYPE_INTEGER = "integer"
TYPE_FLOAT = "float"

# Upload: the two accepted multipart field names, in precedence order.
UPLOAD_FIELDS = ("file", "attachment")
MULTIPART_TYPE = "multipart/form-data"

# Format sniffing. Content decides the format; a filename/URL extension is
# only a hint about whether `charset` should be pre-validated.
ZIP_MAGIC = b"PK\x03\x04"
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
SPREADSHEET_SUFFIXES = (".xls", ".xlsx")
# Signatures that are definitely neither CSV text nor a workbook.
OPAQUE_MAGIC = (
    (b"%PDF-", "PDF"),
    (b"\x89PNG\r\n\x1a\n", "PNG"),
    (b"GIF87a", "GIF"),
    (b"GIF89a", "GIF"),
    (b"\xff\xd8\xff", "JPEG"),
    (b"\x1f\x8b", "gzip"),
    (b"BZh", "bzip2"),
    (b"\xfd7zXZ\x00", "xz"),
    (b"7z\xbc\xaf\x27\x1c", "7-zip"),
    (b"Rar!\x1a\x07", "RAR"),
    (b"\x7fELF", "ELF"),
    (b"SQLite format 3\x00", "SQLite"),
    (b"{\\rtf", "RTF"),
    (b"\x00asm", "WebAssembly"),
    (b"OggS", "Ogg"),
    (b"RIFF", "RIFF"),
    (b"\x25\x21PS", "PostScript"),
)

# Wall-clock budget for the filter/sort/paginate phase of a dataset query.
# Overridable only so the timeout path is testable; see AMBIGUITIES T32.
try:
    QUERY_TIMEOUT_MS = max(0.0, float(os.environ.get("DATAGATE_QUERY_TIMEOUT_MS", "")))
except ValueError:
    QUERY_TIMEOUT_MS = 2000.0

# Strict decimal forms only -- no exponents, no nan/inf, no locale separators.
INT_RE = re.compile(r"^[+-]?[0-9]+$")
DEC_RE = re.compile(r"^[+-]?(?:[0-9]+\.[0-9]*|\.[0-9]+)$")
# Control parameters are strict: plain digits only, no sign, space, or exponent.
UINT_RE = re.compile(r"^[0-9]+$")

class StrictFormDataParser(FormDataParser):
    """Surface multipart parse failures instead of silently yielding nothing."""

    def __init__(self, *args, **kwargs):
        kwargs["silent"] = False
        super().__init__(*args, **kwargs)


class StrictRequest(Request):
    form_data_parser_class = StrictFormDataParser


app = Flask(__name__)
app.request_class = StrictRequest

_STORE: dict[str, dict] = {}
_STORE_LOCK = threading.Lock()

# Built-in defaults; `main` replaces them with the resolved configuration.
CACHE_ENABLED = True
MAX_SOURCE_SIZE = None
ORIGIN_ALLOWLIST = None
REQUIRE_TLS = False
STORAGE_DIR = DEFAULT_STORAGE_DIR


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
class ConfigError(ValueError):
    """Invalid configuration or a config-file read failure: a startup error."""


class Config:
    """The five documented settings, already parsed into their types."""

    __slots__ = SETTING_NAMES

    def __init__(self, **values):
        for name in SETTING_NAMES:
            setattr(self, name, values[name])


def parse_bool(key: str, raw: str | None, default: bool) -> bool:
    """`1/true/yes/on` or `0/false/no/off`, case-insensitive and trimmed."""
    if raw is None:
        return default
    value = raw.strip().lower()
    if value in BOOL_TRUE:
        return True
    if value in BOOL_FALSE:
        return False
    raise ConfigError(
        f"invalid {key}={raw!r}: expected one of "
        f"{', '.join(BOOL_TRUE + BOOL_FALSE)} (case-insensitive)"
    )


def parse_size(raw: str | None) -> int | None:
    """`MAX_SOURCE_SIZE` is a plain non-negative integer count of bytes."""
    if raw is None:
        return None
    value = raw.strip()
    if not UINT_RE.match(value):
        raise ConfigError(
            f"invalid MAX_SOURCE_SIZE={raw!r}: expected a non-negative integer "
            "number of bytes"
        )
    return int(value)


def parse_suffixes(raw: str | None) -> list[str] | None:
    """`ORIGIN_ALLOWLIST` is a comma-separated list of domain suffixes.

    An absent -- or entirely empty -- value means no allowlist at all, which
    the spec says lets every request through.
    """
    if raw is None:
        return None
    suffixes = []
    for part in raw.split(LIST_SEPARATOR):
        suffix = part.strip().strip(".").lower()
        if suffix:
            suffixes.append(suffix)
    return suffixes or None


def parse_storage_dir(raw: str | None) -> str:
    if raw is None or not raw.strip():
        return DEFAULT_STORAGE_DIR
    return os.path.abspath(os.path.expanduser(raw.strip()))


def read_config_file(path: str) -> dict[str, str]:
    """`KEY=VALUE` lines; blank and `#` lines ignored; last assignment wins."""
    try:
        with open(path, "r", encoding="utf-8") as handle:
            text = handle.read()
    except OSError as exc:
        raise ConfigError(
            f"cannot read {CONFIG_ENV} file {path!r}: "
            f"{exc.strerror or exc.__class__.__name__}"
        )
    except UnicodeDecodeError:
        raise ConfigError(f"cannot read {CONFIG_ENV} file {path!r}: not UTF-8 text")

    values: dict[str, str] = {}
    for number, line in enumerate(text.splitlines(), 1):
        stripped = line.strip()
        if not stripped or stripped.startswith(COMMENT_PREFIX):
            continue
        key, separator, value = stripped.partition("=")
        if not separator:
            raise ConfigError(
                f"invalid config at {path}:{number}: expected KEY=VALUE, got {line!r}"
            )
        key = key.strip()
        if not key:
            raise ConfigError(f"invalid config at {path}:{number}: missing key")
        values[key] = value.strip()
    return values


def resolve_config(environ) -> Config:
    """Built-in defaults < `DATAGATE_CONFIG` file < direct environment."""
    raw: dict[str, str] = {}
    path = environ.get(CONFIG_ENV)
    if path is not None and path.strip():
        raw.update(read_config_file(path))
    for name in SETTING_NAMES:  # direct environment variables win outright
        if name in environ:
            raw[name] = environ[name]

    return Config(
        MAX_SOURCE_SIZE=parse_size(raw.get("MAX_SOURCE_SIZE")),
        ORIGIN_ALLOWLIST=parse_suffixes(raw.get("ORIGIN_ALLOWLIST")),
        REQUIRE_TLS=parse_bool("REQUIRE_TLS", raw.get("REQUIRE_TLS"), False),
        STORAGE_DIR=parse_storage_dir(raw.get("STORAGE_DIR")),
        CACHE_ENABLED=parse_bool(CACHE_ENV, raw.get(CACHE_ENV), True),
    )


def apply_config(config: Config) -> None:
    """Install the resolved settings and create `STORAGE_DIR` if missing."""
    global CACHE_ENABLED, MAX_SOURCE_SIZE, ORIGIN_ALLOWLIST, REQUIRE_TLS, STORAGE_DIR
    CACHE_ENABLED = config.CACHE_ENABLED
    MAX_SOURCE_SIZE = config.MAX_SOURCE_SIZE
    ORIGIN_ALLOWLIST = config.ORIGIN_ALLOWLIST
    REQUIRE_TLS = config.REQUIRE_TLS
    STORAGE_DIR = config.STORAGE_DIR
    try:
        os.makedirs(STORAGE_DIR, exist_ok=True)
    except OSError as exc:
        raise ConfigError(
            f"cannot create STORAGE_DIR {STORAGE_DIR!r}: "
            f"{exc.strerror or exc.__class__.__name__}"
        )
    if not os.path.isdir(STORAGE_DIR):
        raise ConfigError(f"cannot create STORAGE_DIR {STORAGE_DIR!r}: not a directory")


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------
class DataGateError(Exception):
    """An error that maps onto a documented status code."""

    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


def error(status: int, message: str):
    response = jsonify({"ok": False, "error": message})
    response.status_code = status
    return response


# ---------------------------------------------------------------------------
# Identity
# ---------------------------------------------------------------------------
def dataset_id(source: str) -> str:
    """Same source URL string -> same id, in this process and any other."""
    return hashlib.sha256(source.encode("utf-8")).hexdigest()[:16]


# ---------------------------------------------------------------------------
# Input validation
# ---------------------------------------------------------------------------
def validate_url(source: str) -> str:
    try:
        parts = urlparse(source)
    except ValueError:
        raise DataGateError(400, f"invalid URL: {source!r}")
    if parts.scheme.lower() not in ("http", "https"):
        raise DataGateError(
            400, "invalid URL: source must be an absolute http:// or https:// URL"
        )
    if not parts.netloc or not parts.hostname:
        raise DataGateError(400, "invalid URL: missing host")
    if any(ch.isspace() for ch in source.strip()):
        raise DataGateError(400, "invalid URL: contains whitespace")
    return source


def validate_charset(name: str) -> str:
    """Reject codec names that do not exist or are not text decoders."""
    try:
        codecs.lookup(name)
        b"".decode(name)
    except (LookupError, TypeError, UnicodeDecodeError, ValueError):
        raise DataGateError(400, f"unsupported or malformed charset: {name!r}")
    return name


# ---------------------------------------------------------------------------
# Maximum source size
# ---------------------------------------------------------------------------
def check_source_size(size: int) -> None:
    """`= limit` is accepted, `> limit` is a 400; an unset limit means no max."""
    if MAX_SOURCE_SIZE is not None and size > MAX_SOURCE_SIZE:
        raise DataGateError(
            400,
            f"source is too large: {size} bytes exceeds the configured "
            f"MAX_SOURCE_SIZE of {MAX_SOURCE_SIZE} bytes",
        )


# ---------------------------------------------------------------------------
# Fetching
# ---------------------------------------------------------------------------
def fetch(url: str) -> bytes:
    try:
        response = requests.get(
            url,
            timeout=FETCH_TIMEOUT,
            stream=True,
            headers={"User-Agent": "datagate/1.0", "Accept": "*/*"},
        )
    except requests.RequestException as exc:
        raise DataGateError(404, f"source unreachable: {exc.__class__.__name__}")

    with response:
        if response.status_code >= 400:
            raise DataGateError(
                404, f"remote HTTP error {response.status_code} for {url}"
            )
        chunks, total = [], 0
        try:
            for chunk in response.iter_content(65536):
                if not chunk:
                    continue
                total += len(chunk)
                check_source_size(total)
                if total > MAX_BYTES:
                    raise DataGateError(400, "source is too large to convert")
                chunks.append(chunk)
        except requests.RequestException as exc:
            raise DataGateError(404, f"source unreachable: {exc.__class__.__name__}")
    return b"".join(chunks)


# ---------------------------------------------------------------------------
# Decoding
# ---------------------------------------------------------------------------
BOMS = (
    (codecs.BOM_UTF32_LE, "utf-32"),
    (codecs.BOM_UTF32_BE, "utf-32"),
    (codecs.BOM_UTF8, "utf-8-sig"),
    (codecs.BOM_UTF16_LE, "utf-16"),
    (codecs.BOM_UTF16_BE, "utf-16"),
)


def decode(data: bytes, charset: str | None) -> str:
    """Decode with the caller's charset, else detect, else latin-1."""
    if charset:
        try:
            return data.decode(charset)
        except (LookupError, TypeError, ValueError) as exc:
            raise DataGateError(
                400, f"unsupported or malformed charset {charset!r}: {exc}"
            )

    for bom, encoding in BOMS:
        if data.startswith(bom):
            try:
                return data.decode(encoding)
            except (UnicodeDecodeError, ValueError):
                break

    try:
        return data.decode("utf-8")  # self-identifying, therefore unambiguous
    except UnicodeDecodeError:
        return data.decode("latin-1")  # ambiguous bytes: the never-failing default


# ---------------------------------------------------------------------------
# Tabular parsing
# ---------------------------------------------------------------------------
def _read_rows(text: str, delimiter: str) -> list[tuple[int, list[str]]]:
    """Return (source line number, cells) for every non-blank record."""
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = []
    start = 1
    for raw in reader:
        cells = [cell.strip() for cell in raw]
        if any(cells):
            rows.append((start, cells))
        start = reader.line_num + 1
    return rows


def _score(rows: list[tuple[int, list[str]]]) -> tuple[float, int] | None:
    """How well does this delimiter explain the file? Higher is better."""
    if len(rows) < 2:
        return None
    width = len(rows[0][1])
    if width < 2:
        return None
    consistent = sum(1 for _, cells in rows if len(cells) == width) / len(rows)
    return (consistent, width)


def looks_non_tabular(text: str) -> str | None:
    """Return a reason if the payload is clearly not a delimited table."""
    if "\x00" in text:
        return "content is binary, not tabular"
    stripped = text.strip()
    if not stripped:
        return "source is empty"
    first = stripped[0]
    if first == "<":
        return "content looks like markup (HTML/XML), not tabular"
    if first in "{[":
        return "content looks like JSON, not tabular"
    return None


def parse_table(text: str) -> tuple[list[str], list[list[str]], list[int]]:
    text = text.lstrip("\ufeff")
    reason = looks_non_tabular(text)
    if reason:
        raise DataGateError(400, reason)

    lines = [line for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        raise DataGateError(
            400, "non-tabular content: need at least one header row and one data row"
        )

    best = None
    for delimiter in DELIMITERS:
        rows = _read_rows(text, delimiter)
        score = _score(rows)
        if score is not None and (best is None or score > best[0]):
            best = (score, rows)

    if best is None:
        # No delimiter present: a single-column table, unless it reads as prose.
        if any(len(line.split()) > MAX_PROSE_TOKENS for line in lines):
            raise DataGateError(400, "non-tabular content: no delimiter found")
        rows = _read_rows(text, "\x01")
        if len(rows) < 2:
            raise DataGateError(
                400,
                "non-tabular content: need at least one header row and one data row",
            )
    else:
        rows = best[1]

    header = rows[0][1]
    width = len(header)
    body, rowids = [], []
    for line_no, row in rows[1:]:
        if len(row) < width:
            row = row + [""] * (width - len(row))
        elif len(row) > width:
            row = row[:width]
        body.append(row)
        rowids.append(line_no)

    if not body:
        raise DataGateError(
            400, "non-tabular content: need at least one header row and one data row"
        )
    return header, body, rowids


# ---------------------------------------------------------------------------
# Type inference
# ---------------------------------------------------------------------------
def infer(value: str):
    """Integers and decimals become JSON numbers; everything else stays text."""
    text = value.strip()
    if not text:
        return text
    if INT_RE.match(text):
        try:
            return int(text)
        except ValueError:
            return text
    if DEC_RE.match(text):
        try:
            number = float(text)
        except ValueError:
            return text
        if number == number and number not in (float("inf"), float("-inf")):
            return number
        return text
    return text


# ---------------------------------------------------------------------------
# Format detection and spreadsheet ingestion
# ---------------------------------------------------------------------------
def detect_format(data: bytes) -> str:
    """Classify a payload as 'xlsx', 'xls', 'opaque' or 'csv' from its bytes."""
    if data.startswith(ZIP_MAGIC):
        return "xlsx"
    if data.startswith(OLE2_MAGIC):
        return "xls"
    for magic, _label in OPAQUE_MAGIC:
        if data.startswith(magic):
            return "opaque"
    return "csv"


def looks_like_spreadsheet_name(name: str) -> bool:
    """Does a filename or URL path claim to be a workbook?"""
    return name.lower().rstrip("/").endswith(SPREADSHEET_SUFFIXES)


def _blank(value) -> bool:
    return value == "" or value is None


def _number(value):
    """Spreadsheets store every number as a float; keep whole ones integral."""
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return str(value)
        if value.is_integer() and abs(value) < 2 ** 53:
            return int(value)
        return value
    return value


def _sheet_text(value) -> str:
    return value if isinstance(value, str) else str(value)


def _xlsx_cell(value):
    """Map an openpyxl cell value onto datagate's stored value space."""
    import datetime as _dt
    import decimal as _decimal

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return _number(value)
    if isinstance(value, _decimal.Decimal):
        return _number(float(value))
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, _dt.timedelta):
        return str(value)
    if isinstance(value, str):
        return infer(value)  # same inference the CSV path applies
    return str(value)


def read_xlsx(data: bytes):
    """Grid of the first worksheet of an .xlsx workbook."""
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - dependency is pinned
        raise DataGateError(400, "unrecognized format: .xlsx support is unavailable")

    book = None
    try:
        book = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
        sheets = book.worksheets
        if not sheets:
            raise DataGateError(400, "the workbook has no worksheets")
        grid = [
            (number, [_xlsx_cell(cell) for cell in row])
            for number, row in enumerate(
                sheets[0].iter_rows(values_only=True), start=1
            )
        ]
    except DataGateError:
        raise
    except Exception:
        raise DataGateError(400, "unrecognized format: not a readable .xlsx workbook")
    finally:
        if book is not None:
            try:
                book.close()
            except Exception:
                pass
    return grid


def read_xls(data: bytes):
    """Grid of the first worksheet of a legacy .xls workbook."""
    try:
        import xlrd
    except ImportError:  # pragma: no cover - dependency is pinned
        raise DataGateError(400, "unrecognized format: .xls support is unavailable")

    try:
        book = xlrd.open_workbook(file_contents=data)
        if book.nsheets < 1:
            raise DataGateError(400, "the workbook has no worksheets")
        sheet = book.sheet_by_index(0)
        grid = []
        for row in range(sheet.nrows):
            cells = [
                _xls_cell(sheet.cell(row, column), book.datemode, xlrd)
                for column in range(sheet.ncols)
            ]
            grid.append((row + 1, cells))
    except DataGateError:
        raise
    except Exception:
        raise DataGateError(400, "unrecognized format: not a readable .xls workbook")
    return grid


def _xls_cell(cell, datemode, xlrd):
    kind, value = cell.ctype, cell.value
    if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return ""
    if kind == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if value else "FALSE"
    if kind == xlrd.XL_CELL_NUMBER:
        return _number(float(value))
    if kind == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(value, datemode).isoformat()
        except Exception:
            return _number(float(value))
    return infer(str(value))


def table_from_grid(grid) -> tuple[list[str], list[list], list[int]]:
    """Turn a (row number, cells) grid into columns/rows/rowids.

    Fully blank rows and trailing blank columns are dropped; the retained row
    numbers stay the worksheet's own 1-based numbering.
    """
    rows = [(number, cells) for number, cells in grid if any(not _blank(c) for c in cells)]
    if len(rows) < 2:
        raise DataGateError(
            400,
            "non-tabular content: the first sheet needs a header row and at "
            "least one data row",
        )

    width = max(len(cells) for _, cells in rows)
    while width > 0 and all(
        len(cells) < width or _blank(cells[width - 1]) for _, cells in rows
    ):
        width -= 1
    if width == 0:
        raise DataGateError(400, "non-tabular content: the first sheet has no columns")

    def pad(cells):
        cells = list(cells[:width])
        return cells + [""] * (width - len(cells))

    header = [_sheet_text(cell) for cell in pad(rows[0][1])]
    body = [pad(cells) for _, cells in rows[1:]]
    rowids = [number for number, _ in rows[1:]]
    return header, body, rowids


def ingest(data: bytes, charset: str | None) -> tuple[list[str], list[list], list[int]]:
    """Parse any accepted source format into columns, typed rows, and rowids.

    `charset` applies -- and is validated -- only on the text CSV branch.
    """
    fmt = detect_format(data)
    if fmt == "xlsx":
        return table_from_grid(read_xlsx(data))
    if fmt == "xls":
        return table_from_grid(read_xls(data))
    if fmt == "opaque":
        raise DataGateError(400, "unrecognized format: the source is not tabular")

    if charset is not None:
        validate_charset(charset)
    text = decode(data, charset)
    columns, rows, rowids = parse_table(text)
    return columns, [[infer(cell) for cell in row] for row in rows], rowids


# ---------------------------------------------------------------------------
# Optional enrichment metadata (AMBIGUITIES T71-T85)
# ---------------------------------------------------------------------------
def enrich_requested() -> bool:
    """On for exactly one `enrich=yes`; every other state keeps enrichment off.

    "Single, exact" is taken literally: a repeated parameter, a different
    spelling, and a bare `enrich` are all just "off", never an error (T71, T72).
    """
    values = request.args.getlist(ENRICH_PARAM)
    return len(values) == 1 and values[0] == ENRICH_VALUE


def _missing(value) -> bool:
    """An empty cell: absent, empty, or whitespace only (T77)."""
    if value is None:
        return True
    return isinstance(value, str) and not value.strip()


def column_type(values: list) -> str:
    """Label a column from its non-missing cells; blanks never demote it (T76)."""
    integers = floats = 0
    for value in values:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            return TYPE_TEXT
        if isinstance(value, int):
            integers += 1
        else:
            floats += 1
    if not integers and not floats:
        return TYPE_TEXT  # nothing to infer from, e.g. an entirely blank column
    if not floats:
        return TYPE_INTEGER
    if not integers:
        return TYPE_FLOAT
    return TYPE_NUMBER  # whole and fractional numbers together (T74)


def column_details(columns: list[str], rows: list[list]) -> dict:
    """Per-column profile, keyed by column name; blanks are counted, not valued."""
    details = {}
    for index, name in enumerate(columns):
        present, missing = [], 0
        for row in rows:
            value = row[index] if index < len(row) else ""
            if _missing(value):
                missing += 1
            else:
                present.append(value)
        details[name] = {
            "type": column_type(present),
            "distinct_count": len(set(present)),  # missing is not a value (T75)
            "missing_count": missing,
        }
    return details


def filetype_label(data: bytes) -> str:
    """The `dataset_summary.filetype` label, decided by content (T73)."""
    if detect_format(data) in SPREADSHEET_FORMATS:
        return FILETYPE_EXCEL
    return FILETYPE_CSV


def build_metadata(filetype: str, table) -> dict:
    """The stored enrichment payload; workbooks get no per-column profile."""
    columns, rows, _rowids = table
    metadata = {
        "dataset_summary": {
            "filetype": filetype,
            "row_count": len(rows),      # data rows, header excluded (T82)
            "column_count": len(columns),
        }
    }
    if filetype != FILETYPE_EXCEL:
        metadata["column_details"] = column_details(columns, rows)
    return metadata


def enrichment_for(data: bytes, table) -> dict:
    """Compute metadata, turning any surprise into a standard JSON error (T83)."""
    try:
        return build_metadata(filetype_label(data), table)
    except DataGateError:
        raise
    except Exception as exc:
        raise DataGateError(500, f"enrichment failed: {exc.__class__.__name__}")


def parse_cache_enabled(raw: str | None) -> bool:
    """Kept as a named entry point for the documented `CACHE_ENABLED` values."""
    return parse_bool(CACHE_ENV, raw, True)


def force_requested() -> bool:
    """`force` is a presence flag: `?force`, at most once, never `?force=...`.

    Flask collapses `?force` and `?force=` to the same empty string, so the raw
    query string is what decides (AMBIGUITIES T51).
    """
    seen = 0
    for part in request.query_string.decode("latin-1").split("&"):
        if not part:
            continue
        name, sep, _ = part.partition("=")
        if unquote_plus(name) != FORCE_PARAM:
            continue
        if sep:
            raise DataGateError(
                400, f"{FORCE_PARAM} is a presence flag and takes no value"
            )
        seen += 1
    if seen > 1:
        raise DataGateError(400, f"repeated query parameter: {FORCE_PARAM}")
    return seen == 1


# ---------------------------------------------------------------------------
# Storage: an in-process index backed by files under STORAGE_DIR, so that a
# dataset stays queryable across a restart that reuses the same directory.
# ---------------------------------------------------------------------------
def dataset_file(identifier: str) -> str:
    return os.path.join(STORAGE_DIR, identifier + DATASET_FILE_SUFFIX)


def _persist(identifier: str, record: dict) -> None:
    """Write the dataset out atomically; a broken store never fails a request."""
    try:
        os.makedirs(STORAGE_DIR, exist_ok=True)
        target = dataset_file(identifier)
        temporary = f"{target}.{os.getpid()}.tmp"
        with open(temporary, "w", encoding="utf-8") as handle:
            json.dump(record, handle)
        os.replace(temporary, target)
    except (OSError, TypeError, ValueError):
        pass


def _restore(identifier: str) -> dict | None:
    """Read a dataset persisted by an earlier run, if the directory still has it."""
    if not DATASET_ID_RE.match(identifier):
        return None  # ids are hex digests; never let a path escape STORAGE_DIR
    try:
        with open(dataset_file(identifier), "r", encoding="utf-8") as handle:
            record = json.load(handle)
    except (OSError, ValueError):
        return None
    if not isinstance(record, dict):
        return None
    if not all(key in record for key in ("columns", "rows", "rowids")):
        return None
    record.setdefault("source", None)
    metadata = record.get("metadata")
    record["metadata"] = metadata if isinstance(metadata, dict) else None
    record["enriched"] = bool(record.get("enriched")) and record["metadata"] is not None
    with _STORE_LOCK:
        _STORE.setdefault(identifier, record)
        return _STORE[identifier]


def load_dataset(identifier: str) -> dict | None:
    with _STORE_LOCK:
        stored = _STORE.get(identifier)
    if stored is not None:
        return stored
    return _restore(identifier)


def endpoint_url(identifier: str) -> str:
    """Relative by default; an absolute `https://` URL when TLS is required."""
    path = f"/datasets/{identifier}"
    if REQUIRE_TLS:
        return f"https://{request.host}{path}"
    return path


def cached_record(identifier: str, source: str) -> dict | None:
    """The already-stored dataset for `source`, if one is cached."""
    stored = load_dataset(identifier)
    if stored is None or stored.get("source") != source:
        return None
    return stored


def store_dataset(identifier: str, source: str, table, metadata: dict | None = None) -> str:
    columns, rows, rowids = table
    record = {
        "source": source,
        "columns": columns,
        "rows": rows,
        "rowids": rowids,
        "enriched": metadata is not None,
        "metadata": metadata,
    }
    with _STORE_LOCK:
        _STORE[identifier] = record
    _persist(identifier, record)
    return endpoint_url(identifier)


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/convert", methods=["GET"])
def convert():
    source = request.args.get("source", "")
    if not source:
        raise DataGateError(400, "missing required query parameter: source")
    validate_url(source)
    forced = force_requested()
    enrich = enrich_requested()

    charset = request.args.get("charset") or None
    # `charset` validates only for text CSV, which is only known after the
    # fetch -- except when the URL itself announces a workbook (AMBIGUITIES T43).
    if charset is not None and not looks_like_spreadsheet_name(urlparse(source).path):
        validate_charset(charset)

    identifier = dataset_id(source)
    if CACHE_ENABLED and not forced:
        stored = cached_record(identifier, source)
        # `enrich=yes` on a cached non-enriched dataset is the one cache hit
        # that still re-ingests, so it can upgrade the stored state; every
        # other hit leaves stored enrichment exactly as it was.
        if stored is not None and (not enrich or stored.get("enriched")):
            # A cache hit is indistinguishable from a fresh parse.
            return jsonify({"ok": True, "endpoint": endpoint_url(identifier)})

    # Re-ingestion only replaces the stored dataset once it -- and its
    # metadata -- have fully succeeded, so a failure leaves any prior dataset
    # queryable with the enrichment it already had.
    data = fetch(source)
    table = ingest(data, charset)
    metadata = enrichment_for(data, table) if enrich else None
    endpoint = store_dataset(identifier, source, table, metadata)
    return jsonify({"ok": True, "endpoint": endpoint})


@app.route("/upload", methods=["POST"])
def upload():
    payload = multipart_payload()
    check_source_size(len(payload))  # the uploaded file's own byte length
    charset = request.args.get("charset") or None
    endpoint = store_dataset(upload_id(payload), None, ingest(payload, charset))
    return jsonify({"ok": True, "endpoint": endpoint})


def upload_id(payload: bytes) -> str:
    """Same file bytes -> same dataset id, mirroring /convert's URL hashing."""
    return hashlib.sha256(payload).hexdigest()[:16]


def multipart_payload() -> bytes:
    """The bytes of the `file` or `attachment` part of a multipart upload."""
    mimetype = (request.mimetype or "").lower()
    if not mimetype.startswith("multipart/"):
        raise DataGateError(
            415,
            "unsupported media type: /upload requires a multipart/form-data body, "
            f"got {mimetype or 'no Content-Type'}",
        )
    if mimetype != MULTIPART_TYPE:
        raise DataGateError(
            400, f"malformed multipart body: expected {MULTIPART_TYPE}, got {mimetype}"
        )
    if not request.mimetype_params.get("boundary"):
        raise DataGateError(400, "malformed multipart body: missing boundary parameter")

    try:
        files, form = request.files, request.form
    except HTTPException:
        raise
    except Exception as exc:
        raise DataGateError(400, f"malformed multipart body: {exc.__class__.__name__}")

    for field in UPLOAD_FIELDS:
        storage = files.get(field)
        if storage is not None:
            return storage.read()
    for field in UPLOAD_FIELDS:  # a same-named text part is still an upload
        if field in form:
            return form[field].encode("utf-8", "surrogateescape")

    if not files and not form:
        raise DataGateError(400, "malformed multipart body: no parts were found")
    raise DataGateError(
        400,
        "missing file field: expected a multipart part named 'file' or 'attachment'",
    )


def _lookup(dataset: str) -> dict:
    stored = load_dataset(dataset)
    if stored is None:
        raise DataGateError(404, f"unknown dataset id: {dataset}")
    return stored


def _plan(stored: dict, started: float):
    """Validate the query controls, then filter -> sort -> paginate.

    Shared by the JSON and CSV views so both honour exactly the same filters,
    sort, and pagination; the CSV view simply ignores the presentation flags.
    """
    raw = {name: _control(name) for name in CONTROL_PARAMS}

    size = _bounded_int("_size", raw["_size"], 1, ROW_LIMIT)
    offset = _bounded_int("_offset", raw["_offset"], 0, 0)
    shape = _shape(raw["_shape"])
    hide_rowid = _toggle("_rowid", raw["_rowid"])
    hide_total = _toggle("_total", raw["_total"])

    columns, rows = stored["columns"], stored["rows"]

    filters = _filters(columns)

    # Both sort parameters are validated even though only one can take effect.
    ascending = _sort_column("_sort", raw["_sort"], columns)
    descending = _sort_column("_sort_desc", raw["_sort_desc"], columns)

    deadline = started + QUERY_TIMEOUT_MS / 1000.0

    order = _apply_filters(rows, filters, deadline)  # spec: filtering precedes sorting
    if descending is not None:  # spec: _sort_desc wins
        order.sort(key=lambda i: _order_key(rows[i][descending]), reverse=True)
    elif ascending is not None:
        order.sort(key=lambda i: _order_key(rows[i][ascending]))
    _check_deadline(deadline)

    window = order[offset : offset + size]  # pagination runs on filtered+sorted rows
    return window, len(order), shape, hide_rowid, hide_total


@app.route("/datasets/<dataset>", methods=["GET"])
def query_dataset(dataset: str):
    started = time.perf_counter()
    stored = _lookup(dataset)
    window, matched, shape, hide_rowid, hide_total = _plan(stored, started)
    columns, rows, rowids = stored["columns"], stored["rows"], stored["rowids"]

    if shape == "objects":
        out = []
        for i in window:
            row = {} if hide_rowid else {"rowid": rowids[i]}
            row.update(zip(columns, rows[i]))
            out.append(row)
    else:
        out = [rows[i] for i in window]

    payload = {"ok": True, "columns": columns, "rows": out}
    if stored.get("enriched") and stored.get("metadata"):
        # `dataset_summary`, plus `column_details` for a CSV dataset. Omitted
        # entirely -- never null or empty -- when the dataset is not enriched.
        payload.update(stored["metadata"])
    if not hide_total:
        payload["total"] = matched  # filtered rows, before pagination
    payload["query_ms"] = max(0.0, round((time.perf_counter() - started) * 1000.0, 3))
    return jsonify(payload)


@app.route("/datasets/<dataset>/export", methods=["GET"])
def export_dataset(dataset: str):
    """The same filtered/sorted/paginated rows, as a CSV attachment.

    `_shape`, `_rowid`, and `_total` are still validated -- they are the same
    query controls -- but a flat CSV has nowhere to express them.
    """
    started = time.perf_counter()
    stored = _lookup(dataset)
    window, _matched, _shape_, _hide_rowid, _hide_total = _plan(stored, started)
    columns, rows = stored["columns"], stored["rows"]

    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(columns)  # source column order, header first
    for index in window:
        writer.writerow([_csv_cell(cell) for cell in rows[index]])

    response = Response(buffer.getvalue().encode("utf-8"), status=200)
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = f'attachment; filename="{dataset}.csv"'
    return response


def _csv_cell(value) -> str:
    """Render a stored value for CSV; numbers keep their JSON spelling."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    return str(value)  # int/float: the same spelling the JSON view uses


def _check_deadline(deadline: float) -> None:
    """A query that outruns its budget is a 400, not a slow success."""
    if time.perf_counter() > deadline:
        raise DataGateError(400, "query timeout: the query took too long to run")


def _filters(columns: list[str]) -> list[tuple[int, str, object]]:
    """Parse `<column>__<comparator>=<value>` params into applicable filters.

    Names beginning with `_` are controls, and names without `__` are not
    filters at all; both are ignored here.
    """
    parsed = []
    for name, values in request.args.lists():
        if name.startswith("_") or FILTER_SEP not in name:
            continue
        if len(values) > 1:
            raise DataGateError(400, f"duplicate filter key: {name}")
        column, _, comparator = name.rpartition(FILTER_SEP)  # column may contain __
        if comparator not in COMPARATORS:
            raise DataGateError(
                400,
                f"invalid comparator {comparator!r} in filter {name!r}: expected one "
                "of exact, contains, less, greater",
            )
        if column not in columns:
            raise DataGateError(400, f"unknown filter column: {column!r}")
        value: object = values[0]
        if comparator in NUMERIC_COMPARATORS:
            number = _numeric(value)
            if number is None:
                raise DataGateError(
                    400,
                    f"invalid {name}: {comparator} needs a numeric value, "
                    f"got {values[0]!r}",
                )
            value = number
        parsed.append((columns.index(column), comparator, value))
    return parsed


def _numeric(value) -> float | None:
    """`float` parse; None means "not numeric", never an exception."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _as_text(value) -> str:
    """The textual form of a stored value, for the string comparators."""
    return value if isinstance(value, str) else str(value)


def _matches(row: list, filters: list[tuple[int, str, object]]) -> bool:
    """Every filter must hold: multiple filters are ANDed."""
    for index, comparator, value in filters:
        cell = row[index]
        if comparator == "exact":
            if _as_text(cell) != value:
                return False
        elif comparator == "contains":
            if value not in _as_text(cell):
                return False
        else:
            number = _numeric(cell)
            if number is None:  # non-numeric stored value: never matched
                return False
            if comparator == "less":
                if not number < value:
                    return False
            elif not number > value:
                return False
    return True


def _apply_filters(rows: list[list], filters, deadline: float) -> list[int]:
    """Indices of the rows that survive every filter, in stored order."""
    _check_deadline(deadline)
    if not filters:
        return list(range(len(rows)))
    kept = []
    for index, row in enumerate(rows):
        if index % ROWS_PER_DEADLINE_CHECK == 0:
            _check_deadline(deadline)
        if _matches(row, filters):
            kept.append(index)
    return kept


def _control(name: str) -> str | None:
    """One value at most: any repeated control parameter is a 400."""
    values = request.args.getlist(name)
    if len(values) > 1:
        raise DataGateError(400, f"repeated control parameter: {name}")
    return values[0] if values else None


def _bounded_int(name: str, raw: str | None, minimum: int, default: int) -> int:
    if raw is None:
        return default
    if not UINT_RE.match(raw) or int(raw) < minimum:
        raise DataGateError(
            400, f"invalid {name}: expected an integer >= {minimum}, got {raw!r}"
        )
    return int(raw)


def _shape(raw: str | None) -> str:
    if raw is None:
        return "lists"
    if raw not in SHAPES:
        raise DataGateError(
            400, f"invalid _shape: expected 'lists' or 'objects', got {raw!r}"
        )
    return raw


def _toggle(name: str, raw: str | None) -> bool:
    """A visibility toggle accepts exactly one value: `hide`."""
    if raw is None:
        return False
    if raw != "hide":
        raise DataGateError(
            400, f"invalid {name}: the only accepted value is 'hide', got {raw!r}"
        )
    return True


def _sort_column(name: str, raw: str | None, columns: list[str]) -> int | None:
    if raw is None:
        return None
    if not raw:
        raise DataGateError(400, f"invalid {name}: a column name is required")
    if raw not in columns:
        raise DataGateError(400, f"unknown column in {name}: {raw!r}")
    return columns.index(raw)


def _order_key(value):
    """A total order over stored values: numbers first, then text."""
    if isinstance(value, bool):
        return (1, 0, str(value))
    if isinstance(value, (int, float)):
        return (0, value, "")
    return (1, 0, str(value))


# ---------------------------------------------------------------------------
# Origin allowlist
#
# Registered before every other `before_request` hook so that the check really
# does run "before routing": an unknown path, a malformed query, and a CORS
# preflight are all refused the same way (AMBIGUITIES T58).
# ---------------------------------------------------------------------------
def referer_hostname(value: str) -> str | None:
    """The hostname of a `Referer`, lowercased and without a trailing dot."""
    try:
        parts = urlparse(value.strip())
    except ValueError:
        return None
    host = parts.hostname
    if not host:
        return None
    return host.rstrip(".").lower()


def host_allowed(host: str) -> bool:
    """Suffix match on whole labels: `example.com` or `*.example.com`."""
    for suffix in ORIGIN_ALLOWLIST or ():
        if host == suffix or host.endswith("." + suffix):
            return True
    return False


@app.before_request
def enforce_origin_allowlist():
    if not ORIGIN_ALLOWLIST:  # no allowlist means all requests pass
        return None
    raw = request.headers.get("Referer")
    if raw is None or not raw.strip():
        raise DataGateError(
            403, "forbidden: this deployment requires a Referer header"
        )
    host = referer_hostname(raw)
    if host is None:
        raise DataGateError(
            403, f"forbidden: Referer {raw!r} has no parseable hostname"
        )
    if not host_allowed(host):
        raise DataGateError(
            403, f"forbidden: origin {host!r} is not in ORIGIN_ALLOWLIST"
        )
    return None


# ---------------------------------------------------------------------------
# Envelope, CORS, and error handling
# ---------------------------------------------------------------------------
@app.after_request
def add_cors_headers(response: Response) -> Response:
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "*"
    response.headers["Access-Control-Expose-Headers"] = "*"
    response.headers["Access-Control-Max-Age"] = "86400"
    return response


@app.before_request
def answer_preflight():
    """Any OPTIONS request is a browser preflight; never let it 404/405."""
    if request.method == "OPTIONS":
        return Response(status=204)
    return None


@app.errorhandler(DataGateError)
def handle_datagate_error(exc: DataGateError):
    return error(exc.status, exc.message)


@app.errorhandler(HTTPException)
def handle_http_exception(exc: HTTPException):
    message = exc.description or exc.name
    response = error(exc.code or 500, message)
    if exc.code == 405 and exc.get_response().headers.get("Allow"):
        response.headers["Allow"] = exc.get_response().headers["Allow"]
    return response


@app.errorhandler(Exception)
def handle_unexpected(exc: Exception):
    return error(500, f"internal error: {exc.__class__.__name__}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="datagate", description=__doc__)
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    parser.add_argument("--address", default=DEFAULT_ADDRESS)
    subparsers = parser.add_subparsers(dest="command")
    start = subparsers.add_parser("start", help="run the HTTP service")
    start.add_argument("--port", type=int, default=argparse.SUPPRESS)
    start.add_argument("--address", default=argparse.SUPPRESS)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv if argv is not None else sys.argv[1:])
    try:
        apply_config(resolve_config(os.environ))
    except ConfigError as exc:
        print(f"datagate: {exc}", file=sys.stderr)
        return 2
    app.run(host=args.address, port=args.port, threaded=True, debug=False)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
