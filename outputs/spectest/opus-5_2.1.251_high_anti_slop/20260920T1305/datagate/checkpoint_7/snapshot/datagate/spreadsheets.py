"""Readers for the spreadsheet formats accepted alongside CSV: ``.xlsx`` and ``.xls``.

Both formats are recognised by their file signature rather than by a file name, and
only the workbook's first worksheet is ingested.
"""

import io
import struct
import zipfile
from datetime import date, datetime, time

import openpyxl
import xlrd

from .errors import DataGateError
from .tables import EXCEL, Table, build_table
from .values import Value

XLSX_SIGNATURE = b"PK\x03\x04"
XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

XLSX_READ_ERRORS = (KeyError, zipfile.BadZipFile)
"""What openpyxl raises for an archive that is not, or no longer, a workbook."""

XLS_READ_ERRORS = (xlrd.XLRDError, xlrd.compdoc.CompDocError, struct.error)
"""The unrelated exception types xlrd raises for a damaged compound document."""


def looks_like_spreadsheet(data: bytes) -> bool:
    """Report whether ``data`` opens with an ``.xlsx`` or ``.xls`` file signature."""
    return data.startswith((XLSX_SIGNATURE, XLS_SIGNATURE))


def read_workbook(data: bytes) -> Table:
    """Parse the first worksheet of a spreadsheet into a ``Table``.

    Any further worksheet is ignored. Bytes that carry a workbook signature but cannot
    be opened as one raise a 400 error, as does a first sheet that is not tabular.
    """
    grid = _read_xlsx(data) if data.startswith(XLSX_SIGNATURE) else _read_xls(data)
    return build_table([_trim(row) for row in grid], EXCEL)


def _trim(row: list[Value]) -> list[Value]:
    """Drop the blank cells a worksheet pads its rows out to the used range with.

    Without this a header row would gain unnamed columns, and every data row would be
    padded to the widest row in the sheet rather than to the header.
    """
    width = len(row)
    while width and row[width - 1] == "":
        width -= 1
    return row[:width]


def _read_xlsx(data: bytes) -> list[list[Value]]:
    """Return the first worksheet of an ``.xlsx`` workbook as a grid of cells."""
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except XLSX_READ_ERRORS as exc:
        raise DataGateError(
            f"source is not a readable .xlsx workbook: {exc}", 400
        ) from exc

    sheet = workbook.worksheets[0]
    return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]


def _read_xls(data: bytes) -> list[list[Value]]:
    """Return the first worksheet of an ``.xls`` workbook as a grid of cells."""
    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except XLS_READ_ERRORS as exc:
        raise DataGateError(
            f"source is not a readable .xls workbook: {exc}", 400
        ) from exc

    sheet = workbook.sheet_by_index(0)
    return [
        [_xls_cell(cell, workbook.datemode) for cell in sheet.row(index)]
        for index in range(sheet.nrows)
    ]


def _xls_cell(cell: xlrd.sheet.Cell, datemode: int) -> Value:
    """Return one ``.xls`` cell, resolving the workbook's date serial numbers first."""
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode).isoformat()
    return _cell(cell.value)


def _cell(value: object) -> Value:
    """Return one worksheet cell as the JSON value the API exposes.

    Dates become ISO-8601 text and whole numbers become integers, so that a worksheet
    and the equivalent CSV describe their rows the same way. Blank cells read as empty
    text, matching how a CSV's missing trailing cells are filled in.
    """
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    return "" if value is None else str(value)
