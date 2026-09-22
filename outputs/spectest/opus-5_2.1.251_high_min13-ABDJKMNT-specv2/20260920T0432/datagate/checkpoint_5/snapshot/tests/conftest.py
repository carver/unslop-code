"""Shared fixtures: a Flask test client and a real local HTTP origin server.

The origin server lets the tests exercise the whole ingestion path (fetch,
decode, parse) against genuine HTTP responses instead of a mocked transport.
"""

import io
import threading
from collections import Counter
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest

from gateway.app import create_app


class _Origin:
    """A tiny HTTP server whose responses the tests register per path."""

    def __init__(self):
        self.responses = {}
        self.hits = Counter()
        origin = self

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                path = self.path.split("?")[0]
                origin.hits[path] += 1
                if path not in origin.responses:
                    self.send_error(404)
                    return
                status, content_type, body = origin.responses[path]
                self.send_response(status)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *args):
                pass

        self._server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self._thread = threading.Thread(target=self._server.serve_forever, daemon=True)
        self._thread.start()

    @property
    def port(self):
        return self._server.server_address[1]

    def serve(self, path, body, status=200, content_type="text/csv"):
        """Register `body` (bytes or str) at `path` and return its full URL."""
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.responses[path] = (status, content_type, body)
        return self.url(path)

    def url(self, path):
        return f"http://127.0.0.1:{self.port}{path}"

    def shutdown(self):
        self._server.shutdown()
        self._server.server_close()


@pytest.fixture(scope="session")
def origin():
    server = _Origin()
    yield server
    server.shutdown()


@pytest.fixture
def client(configured_client):
    return configured_client()


@pytest.fixture
def configured_client(monkeypatch):
    """Build a test client for an app created under a `CACHE_ENABLED` value.

    Passing nothing leaves the variable unset, which is the default
    configuration the rest of the suite runs against.
    """

    def _configured_client(cache_enabled=None):
        monkeypatch.delenv("CACHE_ENABLED", raising=False)
        if cache_enabled is not None:
            monkeypatch.setenv("CACHE_ENABLED", cache_enabled)
        app = create_app()
        app.config.update(TESTING=True)
        return app.test_client()

    return _configured_client


@pytest.fixture
def convert(client, origin):
    """Serve a CSV body and run it through /convert, returning the response."""

    def _convert(body, path=None, params=None, **serve_kwargs):
        path = path or f"/fixture-{abs(hash((body, str(params)))):x}.csv"
        source = origin.serve(path, body, **serve_kwargs)
        query = {"source": source, **(params or {})}
        return client.get("/convert", query_string=query)

    return _convert


@pytest.fixture
def dataset(client, convert):
    """Convert a CSV body and return the parsed /datasets/<id> payload."""

    def _dataset(body, params=None, query=None, **serve_kwargs):
        endpoint = convert(body, params=params, **serve_kwargs).get_json()["endpoint"]
        return client.get(endpoint, query_string=query or {}).get_json()

    return _dataset


@pytest.fixture
def reader(client, convert):
    """Convert a CSV body once, then issue repeated reads of its dataset.

    Returns a function taking the body and giving back `read(query)`, so a test
    can vary control parameters without re-serving and re-converting the CSV.
    """

    def _reader(body, params=None, **serve_kwargs):
        endpoint = convert(body, params=params, **serve_kwargs).get_json()["endpoint"]

        def _read(query=None):
            return client.get(endpoint, query_string=query or {})

        return _read

    return _reader


def _write_xlsx(sheets: dict) -> bytes:
    """An .xlsx workbook holding `sheets` as {title: rows}, first entry first."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title=title)
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_xls(sheets: dict) -> bytes:
    """An .xls workbook holding `sheets` as {title: rows}, first entry first."""
    import xlwt

    workbook = xlwt.Workbook()
    for title, rows in sheets.items():
        sheet = workbook.add_sheet(title)
        for y, row in enumerate(rows):
            for x, value in enumerate(row):
                sheet.write(y, x, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def xlsx():
    """Build .xlsx bytes from rows, or from {sheet title: rows}."""

    def _xlsx(rows_or_sheets):
        return _write_xlsx(_as_sheets(rows_or_sheets))

    return _xlsx


@pytest.fixture
def xls():
    """Build .xls bytes from rows, or from {sheet title: rows}."""

    def _xls(rows_or_sheets):
        return _write_xls(_as_sheets(rows_or_sheets))

    return _xls


def _as_sheets(rows_or_sheets) -> dict:
    if isinstance(rows_or_sheets, dict):
        return rows_or_sheets
    return {"Sheet1": rows_or_sheets}


@pytest.fixture
def upload(client):
    """POST a file to /upload as multipart, returning the response."""

    def _upload(content, field="file", filename="table.csv", params=None):
        if isinstance(content, str):
            content = content.encode("utf-8")
        return client.post(
            "/upload",
            data={field: (io.BytesIO(content), filename)},
            content_type="multipart/form-data",
            query_string=params or {},
        )

    return _upload


@pytest.fixture
def uploaded(client, upload):
    """Upload a file and return the parsed /datasets/<id> payload."""

    def _uploaded(content, query=None, **upload_kwargs):
        endpoint = upload(content, **upload_kwargs).get_json()["endpoint"]
        return client.get(endpoint, query_string=query or {}).get_json()

    return _uploaded
