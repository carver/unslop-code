"""Checks for CSV export, file upload and multi-format ingestion."""

import csv
import io
import socket
import subprocess
import sys
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import quote

import requests

FIXTURES = {}


class Origin(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, *a):
        pass

    def do_GET(self):
        path = self.path.split("?")[0]
        if path not in FIXTURES:
            self.send_response(404)
            self.send_header("Content-Length", "9")
            self.end_headers()
            self.wfile.write(b"not found")
            return
        status, ctype, body = FIXTURES[path]
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


PASS = FAIL = 0


def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  ok   %s" % name)
    else:
        FAIL += 1
        print("  FAIL %s %s" % (name, extra))


def make_xlsx(rows, extra_sheet=None, sheet_title="Sheet1"):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for row in rows:
        ws.append(list(row))
    if extra_sheet is not None:
        ws2 = wb.create_sheet("Second")
        for row in extra_sheet:
            ws2.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_xls(rows, extra_sheet=None):
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("First")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    if extra_sheet is not None:
        sheet2 = book.add_sheet("Second")
        for r, row in enumerate(extra_sheet):
            for c, value in enumerate(row):
                sheet2.write(r, c, value)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def parse_csv(text):
    return [r for r in csv.reader(io.StringIO(text, newline=""))]


def main():
    origin_port = free_port()
    app_port = free_port()

    BASIC = b"name,age,score\nAlice,30,1.5\nBob,25,2\nCarol,41,9\n"
    XLSX = make_xlsx(
        [["name", "age", "score"], ["Alice", 30, 1.5], ["Bob", 25, 2]],
        extra_sheet=[["ignored"], ["nope"]],
    )
    XLS = make_xls(
        [["name", "age", "score"], ["Alice", 30, 1.5], ["Bob", 25, 2]],
        extra_sheet=[["ignored"], ["nope"]],
    )
    XLSX_HEADER_ONLY = make_xlsx([["a", "b", "c"]])
    XLS_HEADER_ONLY = make_xls([["a", "b", "c"]])
    XLSX_EMPTY = make_xlsx([[None]])
    LATIN = "name,city\nJosé,Málaga\n".encode("latin-1")
    UNICODE_CSV = "name,city\nJosé,Málaga\n".encode("utf-8")
    QUOTED = b'name,note\n"Smith, John","said ""hi"", ok"\nPlain,fine\n'
    PDF = b"%PDF-1.4\n\x00\x01\x02\x03binary\x00stuff"
    ZIP_NOT_XLSX = None
    buf = io.BytesIO()
    import zipfile

    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "world")
    ZIP_NOT_XLSX = buf.getvalue()

    FIXTURES["/basic.csv"] = (200, "text/csv", BASIC)
    FIXTURES["/book.xlsx"] = (200, "application/octet-stream", XLSX)
    FIXTURES["/book.xls"] = (200, "application/vnd.ms-excel", XLS)
    FIXTURES["/headeronly.xlsx"] = (200, "application/octet-stream", XLSX_HEADER_ONLY)
    FIXTURES["/headeronly.xls"] = (200, "application/vnd.ms-excel", XLS_HEADER_ONLY)
    FIXTURES["/notxlsx.zip"] = (200, "application/zip", ZIP_NOT_XLSX)
    FIXTURES["/binary"] = (200, "application/pdf", PDF)
    FIXTURES["/latin.csv"] = (200, "text/csv", LATIN)
    FIXTURES["/uni.csv"] = (200, "text/csv", UNICODE_CSV)
    FIXTURES["/quoted.csv"] = (200, "text/csv", QUOTED)
    FIXTURES["/big.csv"] = (200, "text/csv", b"n\n" + b"".join(b"%d\n" % i for i in range(500)))

    origin = ThreadingHTTPServer(("127.0.0.1", origin_port), Origin)
    threading.Thread(target=origin.serve_forever, daemon=True).start()

    proc = subprocess.Popen(
        [sys.executable, "datagate.py", "start", "--port", str(app_port),
         "--address", "127.0.0.1"],
        cwd="/workspace", stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
    )
    base = "http://127.0.0.1:%d" % app_port
    origin_base = "http://127.0.0.1:%d" % origin_port

    for _ in range(100):
        try:
            requests.get(base + "/", timeout=1)
            break
        except Exception:
            time.sleep(0.1)
    else:
        print("server did not start")
        print(proc.stdout.read().decode())
        return 1

    def convert(src, charset=None):
        url = base + "/convert?source=" + quote(src, safe="")
        if charset is not None:
            url += "&charset=" + quote(charset, safe="")
        return requests.get(url, timeout=20)

    def upload(data, field="file", filename="data.csv", charset=None, **kw):
        url = base + "/upload"
        if charset is not None:
            url += "?charset=" + quote(charset, safe="")
        return requests.post(url, files={field: (filename, data)}, timeout=20, **kw)

    def dataset(ep, q=""):
        return requests.get(base + ep + q, timeout=20)

    def export(ep, q=""):
        return requests.get(base + ep + "/export" + q, timeout=20)

    try:
        print("\n[export basics]")
        ep = convert(origin_base + "/basic.csv").json()["endpoint"]
        ds_id = ep.rsplit("/", 1)[-1]
        r = export(ep)
        check("200", r.status_code == 200, r.text)
        check("content-type", r.headers.get("Content-Type") == "text/csv",
              r.headers.get("Content-Type"))
        check("disposition",
              r.headers.get("Content-Disposition") == 'attachment; filename="%s.csv"' % ds_id,
              r.headers.get("Content-Disposition"))
        rows = parse_csv(r.content.decode("utf-8"))
        check("header is source column order", rows[0] == ["name", "age", "score"], rows)
        check("all rows", rows[1:] == [["Alice", "30", "1.5"], ["Bob", "25", "2"],
                                       ["Carol", "41", "9"]], rows)

        print("\n[export honours filters/sort/pagination]")
        rows = parse_csv(export(ep, "?_sort_desc=age").content.decode())
        check("sort desc", [r[0] for r in rows[1:]] == ["Carol", "Alice", "Bob"], rows)
        rows = parse_csv(export(ep, "?_sort=age").content.decode())
        check("sort asc", [r[0] for r in rows[1:]] == ["Bob", "Alice", "Carol"], rows)
        rows = parse_csv(export(ep, "?age__greater=26").content.decode())
        check("filter", [r[0] for r in rows[1:]] == ["Alice", "Carol"], rows)
        rows = parse_csv(export(ep, "?_size=1&_offset=1").content.decode())
        check("paginate", rows[1:] == [["Bob", "25", "2"]], rows)
        rows = parse_csv(export(ep, "?_sort_desc=age&_size=2").content.decode())
        check("filter->sort->paginate order",
              [r[0] for r in rows[1:]] == ["Carol", "Alice"], rows)
        rows = parse_csv(export(ep, "?name__contains=o").content.decode())
        check("contains filter", [r[0] for r in rows[1:]] == ["Bob", "Carol"], rows)

        print("\n[export ignores _shape/_rowid/_total]")
        plain = export(ep).content
        for q in ["?_shape=objects", "?_shape=lists", "?_rowid=hide", "?_total=hide",
                  "?_shape=objects&_rowid=hide&_total=hide"]:
            check("ignored %s" % q, export(ep, q).content == plain,
                  export(ep, q).content[:120])

        print("\n[export defaults + edge cases]")
        big = convert(origin_base + "/big.csv").json()["endpoint"]
        rows = parse_csv(export(big).content.decode())
        check("default page size 100", len(rows) - 1 == 100, len(rows))
        qep = convert(origin_base + "/quoted.csv").json()["endpoint"]
        rows = parse_csv(export(qep).content.decode())
        check("quoting round-trips",
              rows == [["name", "note"], ["Smith, John", 'said "hi", ok'],
                       ["Plain", "fine"]], rows)
        uep = convert(origin_base + "/uni.csv").json()["endpoint"]
        check("unicode export",
              parse_csv(export(uep).content.decode("utf-8"))[1] == ["José", "Málaga"],
              export(uep).content)
        r = export(ep, "?age__greater=999")
        check("empty result still has header",
              parse_csv(r.content.decode()) == [["name", "age", "score"]], r.content)
        r = export("/datasets/deadbeefdeadbeef")
        check("unknown dataset -> 404", r.status_code == 404, r.text)
        check("unknown dataset json", r.json().get("ok") is False, r.text)
        check("unknown dataset json ct",
              "application/json" in r.headers.get("Content-Type", ""))
        r = export(ep, "?_size=0")
        check("bad control -> 400", r.status_code == 400, r.text)
        check("bad control json", r.json().get("ok") is False, r.text)
        r = export(ep, "?_sort=nope")
        check("unknown sort col -> 400", r.status_code == 400, r.text)
        r = export(ep, "?nope__exact=1")
        check("unknown filter col -> 400", r.status_code == 400, r.text)
        check("export has CORS",
              export(ep).headers.get("Access-Control-Allow-Origin") == "*")

        print("\n[upload]")
        r = upload(BASIC)
        check("200", r.status_code == 200, r.text)
        j = r.json()
        check("ok true", j.get("ok") is True, j)
        check("endpoint shape", j.get("endpoint", "").startswith("/datasets/"), j)
        check("only ok+endpoint", set(j) == {"ok", "endpoint"}, j)
        dj = dataset(j["endpoint"]).json()
        check("uploaded rows",
              dj["columns"] == ["name", "age", "score"]
              and dj["rows"][0] == ["Alice", 30, 1.5], dj)
        r2 = upload(BASIC)
        check("same bytes -> same id", r2.json()["endpoint"] == j["endpoint"],
              (r2.json(), j))
        r3 = upload(BASIC, filename="other-name.csv")
        check("same bytes different filename -> same id",
              r3.json()["endpoint"] == j["endpoint"])
        r4 = upload(BASIC, field="attachment")
        check("attachment field works", r4.status_code == 200, r4.text)
        check("attachment same id", r4.json()["endpoint"] == j["endpoint"])
        r5 = upload(b"a,b\n1,2\n")
        check("different bytes -> different id", r5.json()["endpoint"] != j["endpoint"])
        check("export from upload works",
              export(j["endpoint"]).status_code == 200)

        print("\n[upload charset]")
        r = upload(LATIN, charset="latin-1")
        check("explicit latin-1", r.status_code == 200, r.text)
        dj = dataset(r.json()["endpoint"]).json()
        check("latin-1 decoded", dj["rows"] == [["José", "Málaga"]], dj)
        r = upload(LATIN)
        check("charset auto-detected", r.status_code == 200, r.text)
        dj = dataset(r.json()["endpoint"]).json()
        check("fallback decoded", dj["rows"] == [["José", "Málaga"]], dj)
        r = upload(BASIC, charset="no-such-charset")
        check("bad charset on csv -> 400", r.status_code == 400, r.text)
        r = upload(LATIN, charset="utf-8")
        check("undecodable charset -> 400", r.status_code == 400, r.text)
        r = upload(BASIC, charset="")
        check("empty charset -> 400", r.status_code == 400, r.text)

        print("\n[upload errors]")
        r = requests.post(base + "/upload", data=b"name,age\nA,1\n",
                          headers={"Content-Type": "text/csv"}, timeout=10)
        check("non-multipart -> 415", r.status_code == 415, r.text)
        check("415 envelope",
              r.json().get("ok") is False and isinstance(r.json().get("error"), str),
              r.text)
        check("415 json ct", "application/json" in r.headers.get("Content-Type", ""))
        r = requests.post(base + "/upload", json={"a": 1}, timeout=10)
        check("json body -> 415", r.status_code == 415, r.text)
        r = requests.post(base + "/upload", timeout=10)
        check("no body/content-type -> 415", r.status_code == 415, r.text)
        r = requests.post(base + "/upload", data={"a": "b"}, timeout=10)
        check("urlencoded form -> 415", r.status_code == 415, r.text)
        r = requests.post(base + "/upload", files={"other": ("x.csv", BASIC)}, timeout=10)
        check("wrong field -> 400", r.status_code == 400, r.text)
        check("400 envelope", r.json().get("ok") is False, r.text)
        r = requests.post(
            base + "/upload",
            data=b"--boundary\r\ngarbage without headers\r\n",
            headers={"Content-Type": "multipart/form-data; boundary=boundary"},
            timeout=10)
        check("malformed multipart -> 400", r.status_code == 400, r.text)
        r = requests.post(
            base + "/upload", data=b"whatever",
            headers={"Content-Type": "multipart/form-data"}, timeout=10)
        check("multipart without boundary -> 400", r.status_code == 400, r.text)
        r = upload(PDF, filename="x.pdf")
        check("unsupported format upload -> 400", r.status_code == 400, r.text)
        r = upload(b"", filename="empty.csv")
        check("empty upload -> 400", r.status_code == 400, r.text)
        r = upload(ZIP_NOT_XLSX, filename="x.zip")
        check("zip that is not xlsx -> 400", r.status_code == 400, r.text)
        r = requests.get(base + "/upload", timeout=10)
        check("GET /upload not 200", r.status_code != 200, r.status_code)
        check("GET /upload json", "application/json" in r.headers.get("Content-Type", ""))

        print("\n[xlsx ingestion]")
        r = upload(XLSX, filename="book.xlsx")
        check("upload xlsx 200", r.status_code == 200, r.text)
        dj = dataset(r.json()["endpoint"]).json()
        check("xlsx columns", dj["columns"] == ["name", "age", "score"], dj)
        check("xlsx rows", dj["rows"] == [["Alice", 30, 1.5], ["Bob", 25, 2]], dj)
        check("xlsx only first sheet", dj["columns"] != ["ignored"], dj)
        rows = parse_csv(export(r.json()["endpoint"]).content.decode())
        check("xlsx export", rows == [["name", "age", "score"],
                                      ["Alice", "30", "1.5"], ["Bob", "25", "2"]], rows)
        r = convert(origin_base + "/book.xlsx")
        check("convert xlsx 200", r.status_code == 200, r.text)
        dj = dataset(r.json()["endpoint"]).json()
        check("convert xlsx rows", dj["rows"] == [["Alice", 30, 1.5], ["Bob", 25, 2]], dj)
        r = upload(XLSX, filename="b.xlsx", charset="no-such-charset")
        check("charset ignored for xlsx", r.status_code == 200, r.text)
        r = convert(origin_base + "/headeronly.xlsx")
        check("xlsx header only -> 400", r.status_code == 400, r.text)
        r = upload(XLSX_EMPTY, filename="e.xlsx")
        check("xlsx empty sheet -> 400", r.status_code == 400, r.text)
        r = convert(origin_base + "/notxlsx.zip")
        check("convert non-xlsx zip -> 400", r.status_code == 400, r.text)

        print("\n[xls ingestion]")
        r = upload(XLS, filename="book.xls")
        check("upload xls 200", r.status_code == 200, r.text)
        dj = dataset(r.json()["endpoint"]).json()
        check("xls columns", dj["columns"] == ["name", "age", "score"], dj)
        check("xls rows", dj["rows"] == [["Alice", 30, 1.5], ["Bob", 25, 2]], dj)
        rows = parse_csv(export(r.json()["endpoint"]).content.decode())
        check("xls export", rows == [["name", "age", "score"],
                                     ["Alice", "30", "1.5"], ["Bob", "25", "2"]], rows)
        r = convert(origin_base + "/book.xls")
        check("convert xls 200", r.status_code == 200, r.text)
        dj = dataset(r.json()["endpoint"]).json()
        check("convert xls rows", dj["rows"] == [["Alice", 30, 1.5], ["Bob", 25, 2]], dj)
        r = upload(XLS, filename="b.xls", charset="no-such-charset")
        check("charset ignored for xls", r.status_code == 200, r.text)
        r = convert(origin_base + "/headeronly.xls")
        check("xls header only -> 400", r.status_code == 400, r.text)
        check("xls same bytes -> same id",
              upload(XLS, filename="z.xls").json()["endpoint"]
              == upload(XLS, filename="q.xls").json()["endpoint"])

        print("\n[spreadsheet queries]")
        ep_x = upload(XLSX, filename="book.xlsx").json()["endpoint"]
        dj = dataset(ep_x, "?_shape=objects").json()
        check("objects shape works",
              dj["rows"][0] == {"rowid": 2, "name": "Alice", "age": 30, "score": 1.5}, dj)
        dj = dataset(ep_x, "?age__less=28").json()
        check("filter on sheet", dj["rows"] == [["Bob", 25, 2]], dj)
        rows = parse_csv(export(ep_x, "?_sort_desc=age").content.decode())
        check("sorted sheet export", [r[0] for r in rows[1:]] == ["Alice", "Bob"], rows)

        print("\n[multipart parts without a file name]")
        for field in ("file", "attachment"):
            body = (b"--X\r\nContent-Disposition: form-data; name=\"%s\"\r\n\r\n"
                    % field.encode()) + LATIN + b"\r\n--X--\r\n"
            r = requests.post(
                base + "/upload?charset=latin-1", data=body,
                headers={"Content-Type": "multipart/form-data; boundary=X"},
                timeout=10)
            check("bare %s part 200" % field, r.status_code == 200, r.text)
            check("bare %s part keeps exact bytes" % field,
                  dataset(r.json()["endpoint"]).json()["rows"] == [["José", "Málaga"]],
                  r.text)

        print("\n[cors / preflight for upload]")
        r = requests.options(base + "/upload", timeout=10,
                             headers={"Origin": "http://example.com",
                                      "Access-Control-Request-Method": "POST"})
        check("preflight ok", r.status_code < 400, r.status_code)
        check("preflight allows POST",
              "POST" in (r.headers.get("Access-Control-Allow-Methods") or ""),
              dict(r.headers))
        check("upload ACAO",
              upload(BASIC).headers.get("Access-Control-Allow-Origin") == "*")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=5)
        except Exception:
            proc.kill()
        origin.shutdown()

    print("\n%d passed, %d failed" % (PASS, FAIL))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
