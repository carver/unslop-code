"""Reading the first worksheet of an `.xls` or `.xlsx` payload.

Only the first sheet in tab order is ingested (T47). Cells are rendered to text
and then left to the same per-cell inference CSV uses (T44), so a table means the
same thing whichever of the three formats it arrived in -- which matters most for
`.xls`, where every number is stored as a double.
"""

import datetime
from io import BytesIO
from zipfile import BadZipFile

import openpyxl
import xlrd
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.compdoc import CompDocError

from .errors import DataGateError
from .formats import SourceFormat

BOOLEAN_TEXT = {True: "TRUE", False: "FALSE"}

# A container that announces itself as a workbook but that no reader can open is
# an unrecognized format (T43). `LookupError` covers both a ZIP whose workbook
# parts are missing and a book that turns out to hold no worksheets at all.
_UNREADABLE = (BadZipFile, InvalidFileException, LookupError, xlrd.XLRDError, CompDocError)


def read_workbook(data: bytes, source_format: SourceFormat) -> list[list[str]]:
    """Return the first worksheet of `data` as a grid of text cells."""
    read = _read_xlsx if source_format is SourceFormat.XLSX else _read_xls
    try:
        grid = read(data)
    except _UNREADABLE as exc:
        raise DataGateError(f"Unrecognized {source_format.value} file: {exc}", 400) from None
    width = _used_width(grid)
    return [row[:width] for row in grid]


def _read_xlsx(data: bytes) -> list[list[str]]:
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_xls(data: bytes) -> list[list[str]]:
    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    return [
        [_xls_cell_text(sheet.cell(row, column), book.datemode) for column in range(sheet.ncols)]
        for row in range(sheet.nrows)
    ]


def _used_width(grid: list[list[str]]) -> int:
    """How many columns reach the last one that holds anything (T45)."""
    return max((index + 1 for row in grid for index, cell in enumerate(row) if cell.strip()), default=0)


def _cell_text(value: object) -> str:
    """Render one cell as the text the CSV rules would have been given."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return BOOLEAN_TEXT[value]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return _timestamp_text(value)
    return str(value)


def _xls_cell_text(cell: xlrd.sheet.Cell, datemode: int) -> str:
    """`.xls` keeps dates, booleans and errors as bare numbers, so widen them first."""
    if cell.ctype == xlrd.XL_CELL_DATE:
        return _timestamp_text(xlrd.xldate_as_datetime(cell.value, datemode))
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return BOOLEAN_TEXT[bool(cell.value)]
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return ""
    return _cell_text(cell.value)


def _timestamp_text(value: datetime.date | datetime.time) -> str:
    """ISO 8601, dropping a midnight time so date-only cells read as plain dates."""
    if isinstance(value, datetime.datetime) and value.time() == datetime.time():
        return value.date().isoformat()
    return value.isoformat()
